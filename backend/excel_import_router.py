import io
import re
import unicodedata
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, File, Header, HTTPException, UploadFile
from pymongo import UpdateOne
from pymongo.errors import BulkWriteError


ELECTRONIC_INHALATION_CATEGORY = "PRODOTTI DA INALAZIONE SENZA COMBUSTIONE ELETTRONICA"
ELECTRONIC_INHALATION_BRANDS = (
    "ELFBAR", "ELFLIQ", "LOST MARY", "KIWI", "ELFA", "ELFX", "VAPORESSO",
)


def is_electronic_inhalation(description: str) -> bool:
    normalized = " ".join((description or "").upper().split())
    return any(brand in normalized for brand in ELECTRONIC_INHALATION_BRANDS)


def electronic_inhalation_filter() -> Dict[str, Any]:
    return {
        "$or": [
            {"descrizione": {"$regex": re.escape(brand), "$options": "i"}}
            for brand in ELECTRONIC_INHALATION_BRANDS
        ]
    }


def category_from_description(description: str, code: str) -> str:
    normalized = (description or "").upper()
    if is_electronic_inhalation(normalized):
        return ELECTRONIC_INHALATION_CATEGORY
    for keyword in ("VAPORESSO", "POD", "KIT", "MG/ML", "LIQUID", "TEREA"):
        if keyword in normalized:
            return "PRODOTTI DA INALAZIONE SENZA COMBUSTIONE"
    if str(code or "").startswith("AMMS"):
        return "SIGARETTE"
    return "ACCESSORI"


def create_excel_import_router(
    *,
    get_database,
    read_capped,
    create_backup_snapshot,
    restore_backup_collections,
    record_import_history,
    import_smart_venue,
    sync_product_vending_stock,
    product_model,
    historical_order_model,
    physical_shop_stock_from_excel,
    product_code_text,
    canonical_product_code,
    adm_alias_index,
    remap_code_references,
    normalize_nonnegative_money,
    normalize_money,
    decimal_money,
    apply_electronic_categories,
    excel_import_backup_collections,
):
    router = APIRouter()
    import_progress: Dict[str, Dict[str, Any]] = {}

    def set_import_progress(
        job_id: Optional[str],
        status: str,
        message: str,
        current: int = 0,
        total: int = 1,
        **extra: Any,
    ) -> None:
        if not job_id:
            return
        safe_total = max(1, total)
        import_progress[job_id] = {
            "job_id": job_id,
            "status": status,
            "message": message,
            "current": current,
            "total": safe_total,
            "percent": min(100, max(0, round(current / safe_total * 100))),
            "updated_at": datetime.now(timezone.utc).isoformat(),
            **extra,
        }

    async def bulk_upsert(collection, operations, batch_size: int = 1000) -> Dict[str, int]:
        inserted = updated = errors = 0
        for start in range(0, len(operations), batch_size):
            batch = operations[start:start + batch_size]
            try:
                result = await collection.bulk_write(batch, ordered=False)
                inserted += result.upserted_count
                updated += result.matched_count
            except BulkWriteError as exc:
                details = exc.details or {}
                inserted += details.get("nUpserted", 0)
                updated += details.get("nMatched", 0)
                errors += len(details.get("writeErrors", [])) or len(batch)
        return {"inseriti": inserted, "aggiornati": updated, "errori": errors}

    async def bulk_update(collection, operations, batch_size: int = 1000) -> Dict[str, int]:
        matched = modified = errors = 0
        for start in range(0, len(operations), batch_size):
            batch = operations[start:start + batch_size]
            try:
                result = await collection.bulk_write(batch, ordered=False)
                matched += result.matched_count
                modified += result.modified_count
            except BulkWriteError as exc:
                details = exc.details or {}
                matched += details.get("nMatched", 0)
                modified += details.get("nModified", 0)
                errors += len(details.get("writeErrors", [])) or len(batch)
        return {"trovati": matched, "modificati": modified, "errori": errors}

    async def apply_electronic_inhalation_categories() -> int:
        result = await get_database().prodotti.update_many(
            electronic_inhalation_filter(),
            {"$set": {"categoria": ELECTRONIC_INHALATION_CATEGORY}},
        )
        return int(result.modified_count or 0)

    async def import_products(
        worksheet,
        adm_aliases: Optional[Dict[str, List[Dict[str, str]]]] = None,
    ) -> Dict[str, int]:
        database = get_database()
        adm_aliases = adm_aliases or {}
        import_id = str(uuid.uuid4())
        imported_codes = set()
        alias_codes = set()
        alias_mappings: Dict[str, str] = {}
        products_by_code: Dict[str, Dict[str, Any]] = {}
        canonical_rows = set()
        shop_sold_value = Decimal(0)
        vending_sold_value = Decimal(0)
        errors = 0
        for row in worksheet.iter_rows(min_row=3, values_only=True):
            try:
                code = row[0]
                if not code:
                    continue
                raw_code = product_code_text(code)
                description = str(row[1] or "").strip()
                canonical_code = canonical_product_code(raw_code, description, adm_aliases)
                if canonical_code != raw_code:
                    alias_codes.add(raw_code)
                    alias_mappings[raw_code] = canonical_code
                price = normalize_nonnegative_money(row[7] or 0)
                shop_sold = int(row[3] or 0)
                vending_sold = int(row[17] or 0)
                shop_sold_value += shop_sold * decimal_money(price)
                vending_sold_value += vending_sold * decimal_money(price)
                imported_codes.add(canonical_code)
                row_data = {
                    "codice": canonical_code,
                    "descrizione": description,
                    "categoria": category_from_description(description, canonical_code),
                    "acquistati": int(row[2] or 0),
                    "venduti_negozio": shop_sold,
                    "prezzo": price,
                    "giacenza_vending": int(row[13] or 0),
                    "venduti_vending": vending_sold,
                    "presente_ultimo_import": True,
                    "ultimo_import_id": import_id,
                }
                row_data["giacenza_negozio"] = physical_shop_stock_from_excel(
                    row[6], row_data["giacenza_vending"], row_data["venduti_vending"]
                )
                existing = products_by_code.get(canonical_code)
                if existing:
                    if raw_code == canonical_code:
                        products_by_code[canonical_code] = row_data
                        canonical_rows.add(canonical_code)
                    elif canonical_code not in canonical_rows and not existing.get("descrizione"):
                        products_by_code[canonical_code] = row_data
                else:
                    products_by_code[canonical_code] = row_data
                    if raw_code == canonical_code:
                        canonical_rows.add(canonical_code)
            except Exception:
                errors += 1

        operations = []
        for code, data in products_by_code.items():
            insert_defaults = product_model(codice=code, descrizione=data["descrizione"]).model_dump()
            insert_defaults = {key: value for key, value in insert_defaults.items() if key not in data}
            operations.append(UpdateOne(
                {"codice": code},
                {"$set": data, "$setOnInsert": insert_defaults},
                upsert=True,
            ))
        result = await bulk_upsert(database.prodotti, operations)
        if alias_codes:
            for old_code, new_code in alias_mappings.items():
                await remap_code_references(old_code, new_code)
            await database.prodotti.delete_many({"codice": {"$in": sorted(alias_codes)}})
        if imported_codes:
            await database.prodotti.update_many(
                {"codice": {"$nin": list(imported_codes)}},
                {"$set": {"presente_ultimo_import": False}},
            )
        result["errori"] += errors
        result["elettroniche"] = await apply_electronic_categories()
        result["alias_unificati"] = len(alias_codes)
        result["ultimo_import_id"] = import_id
        result["valore_venduto_negozio_excel"] = normalize_money(shop_sold_value)
        result["valore_venduto_vending_excel"] = normalize_money(vending_sold_value)
        result["valore_venduto_totale_excel"] = normalize_money(shop_sold_value + vending_sold_value)
        return result

    def excel_header_key(value: Any) -> str:
        text = unicodedata.normalize("NFKD", str(value or ""))
        ascii_text = text.encode("ascii", "ignore").decode().lower()
        return re.sub(r"[^a-z0-9]", "", ascii_text)

    def find_product_stock_sheet(workbook):
        aliases = {
            "codice": {"codice", "cod", "code"},
            "descrizione": {"descrizione", "descr", "prodotto"},
            "acquistati": {"acquistati", "acquisti"},
            "giacenza_negozio": {"giacenzanegozio", "giacnegozio", "rimanenzenegozio"},
            "giacenza_vending": {"giacenzavending", "giacvending", "rimanenzevending"},
        }
        required = set(aliases)
        for worksheet in workbook.worksheets:
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 10), values_only=True),
                start=1,
            ):
                columns: Dict[str, int] = {}
                for index, value in enumerate(row):
                    key = excel_header_key(value)
                    for field, names in aliases.items():
                        if key in names and field not in columns:
                            columns[field] = index
                if required.issubset(columns):
                    return worksheet, row_number, columns
        return None

    async def import_product_stock_sheet(
        worksheet,
        header_row: int,
        columns: Dict[str, int],
        adm_aliases: Optional[Dict[str, List[Dict[str, str]]]] = None,
    ) -> Dict[str, int]:
        adm_aliases = adm_aliases or {}
        rows_by_code: Dict[str, Dict[str, Any]] = {}
        errors = 0
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            try:
                raw_code = row[columns["codice"]] if len(row) > columns["codice"] else None
                if raw_code in (None, ""):
                    continue
                description = str(row[columns["descrizione"]] or "").strip()
                code = canonical_product_code(raw_code, description, adm_aliases)
                if not code:
                    continue
                rows_by_code[code] = {
                    "codice": code,
                    "descrizione": description,
                    "giacenza_negozio": int(row[columns["giacenza_negozio"]] or 0),
                    "giacenza_vending": int(row[columns["giacenza_vending"]] or 0),
                }
            except (IndexError, TypeError, ValueError):
                errors += 1

        operations = []
        for code, data in rows_by_code.items():
            insert_defaults = product_model(codice=code, descrizione=data["descrizione"]).model_dump()
            insert_defaults = {key: value for key, value in insert_defaults.items() if key not in data}
            operations.append(UpdateOne(
                {"codice": code},
                {"$set": data, "$setOnInsert": insert_defaults},
                upsert=True,
            ))
        result = await bulk_upsert(get_database().prodotti, operations)
        result["errori"] += errors
        result["righe_lette"] = len(rows_by_code)
        return result

    async def import_price_list(worksheet) -> Dict[str, int]:
        operations = []
        errors = 0
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            try:
                if not row[0]:
                    continue
                code = str(row[0]).strip()
                data = {
                    "codice": code,
                    "descrizione": str(row[1] or "").strip(),
                    "prezzo": normalize_nonnegative_money(row[2] or 0),
                    "confezione": str(row[3] or "").strip(),
                }
                operations.append(UpdateOne({"codice": code}, {"$set": data}, upsert=True))
            except Exception:
                errors += 1
        result = await bulk_upsert(get_database().listino_adm, operations)
        result["errori"] += errors
        return result

    async def import_vending(worksheet) -> Dict[str, int]:
        operations = []
        imported_codes = set()
        errors = 0
        snapshot_at = datetime.now(timezone.utc).isoformat()
        for row in worksheet.iter_rows(min_row=4, values_only=True):
            try:
                if not row[0]:
                    continue
                column = str(row[0]).strip()
                if not re.match(r"^[A-Z]\d+$", column):
                    continue
                data = {
                    "colonna": column,
                    "codice": str(row[1] or "").strip(),
                    "descrizione": str(row[2] or "").strip(),
                    "giacenza": int(row[3]) if row[3] is not None else 0,
                    "capacita_max": int(row[4] or 5),
                    "soglia_minima": int(row[5]) if row[5] is not None else 2,
                    "giacenza_aggiornata_il": snapshot_at,
                    "giacenza_sorgente": "IMPORT_EXCEL",
                }
                operations.append(UpdateOne({"colonna": column}, {"$set": data}, upsert=True))
                if data["codice"]:
                    imported_codes.add(data["codice"])
            except Exception:
                errors += 1
        result = await bulk_upsert(get_database().vending, operations)
        for code in sorted(imported_codes):
            await sync_product_vending_stock(code)
        result["prodotti_riconciliati"] = len(imported_codes)
        result["errori"] += errors
        return result

    async def import_order_history(worksheet) -> Dict[str, int]:
        database = get_database()
        inserted = errors = 0
        await database.storico_ordini.delete_many({})
        batch = []
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            try:
                if not row[0]:
                    continue
                date_value = row[0]
                if isinstance(date_value, datetime):
                    date_value = date_value.isoformat()
                order = historical_order_model(
                    data=str(date_value),
                    file_sorgente=str(row[1] or ""),
                    codice=str(row[2] or "").strip(),
                    descrizione=str(row[3] or "").strip(),
                    quantita=int(row[4] or 0),
                    prezzo=normalize_nonnegative_money(row[5] or 0),
                )
                batch.append(order.model_dump())
                if len(batch) >= 1000:
                    await database.storico_ordini.insert_many(batch)
                    inserted += len(batch)
                    batch = []
            except Exception:
                errors += 1
        if batch:
            await database.storico_ordini.insert_many(batch)
            inserted += len(batch)
        return {"inseriti": inserted, "aggiornati": 0, "errori": errors}

    async def import_parameters(worksheet) -> Dict[str, int]:
        operations = []
        skipped = errors = 0
        database = get_database()
        existing = {row["nome"] async for row in database.parametri.find({}, {"_id": 0, "nome": 1})}
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            try:
                name = row[0]
                if not name:
                    continue
                name = str(name).strip()
                value = row[1]
                if value is None:
                    continue
                normalized_value = float(str(value).replace(",", "."))
                if name in existing:
                    operations.append(UpdateOne({"nome": name}, {"$set": {"valore": normalized_value}}))
                else:
                    skipped += 1
            except Exception:
                errors += 1
        result = await bulk_upsert(database.parametri, operations)
        result["saltati"] = skipped
        result["errori"] += errors
        return result

    async def import_historical_sales(worksheet) -> Dict[str, int]:
        aliases = {
            "data": {"data", "giorno"},
            "codice": {"codice", "cod", "code"},
            "descrizione": {"descrizione", "descr", "prodotto"},
            "quantita": {"qta", "quantita", "pezzi"},
            "importo": {"importo", "totale", "valore"},
            "categoria": {"categoria", "cat"},
        }
        columns: Dict[str, int] = {}
        header_row = 0
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 10), values_only=True),
            start=1,
        ):
            candidate: Dict[str, int] = {}
            for index, value in enumerate(row):
                key = excel_header_key(value)
                for field, names in aliases.items():
                    if key in names and field not in candidate:
                        candidate[field] = index
            if {"data", "codice", "quantita"}.issubset(candidate):
                columns = candidate
                header_row = row_number
                break
        if not columns:
            raise HTTPException(
                422,
                "DB_STORICO_VEND: intestazioni Data, Codice e Quantità non riconosciute; storico precedente conservato",
            )

        def value_at(row: tuple, field: str, default: Any = None) -> Any:
            index = columns.get(field)
            return row[index] if index is not None and index < len(row) else default

        def iso_date(value: Any) -> str:
            if isinstance(value, datetime):
                return value.isoformat()
            text = str(value or "").strip()
            for date_format in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%y"):
                try:
                    return datetime.strptime(text, date_format).isoformat()
                except ValueError:
                    pass
            return text

        def product_code(value: Any) -> str:
            if isinstance(value, datetime) and value.year < 1950:
                from openpyxl.utils.datetime import to_excel
                return str(int(to_excel(value)))
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value or "").strip()

        errors = 0
        documents = []
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            try:
                date_value = value_at(row, "data")
                code = product_code(value_at(row, "codice"))
                if not date_value or not code:
                    continue
                documents.append({
                    "id": str(uuid.uuid4()),
                    "data": iso_date(date_value),
                    "codice": code,
                    "descrizione": str(value_at(row, "descrizione", "") or "").strip(),
                    "quantita": int(value_at(row, "quantita", 0) or 0),
                    "importo": normalize_money(value_at(row, "importo", 0) or 0),
                    "categoria": str(value_at(row, "categoria", "") or "").strip(),
                })
            except Exception:
                errors += 1
        collection = get_database().db_storico_vend
        await collection.delete_many({})
        for start in range(0, len(documents), 1000):
            await collection.insert_many(documents[start:start + 1000])
        return {"inseriti": len(documents), "aggiornati": 0, "errori": errors}

    async def import_extended_vending_history(worksheet) -> Dict[str, int]:
        collection = get_database().db_storico_vending_ext
        inserted = errors = 0
        await collection.delete_many({})
        batch = []
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            try:
                if not row or all(cell in (None, "") for cell in row):
                    continue
                batch.append({
                    "id": str(uuid.uuid4()),
                    "raw": [str(cell) if cell is not None else "" for cell in row[:18]],
                })
                if len(batch) >= 1000:
                    await collection.insert_many(batch)
                    inserted += len(batch)
                    batch = []
            except Exception:
                errors += 1
        if batch:
            await collection.insert_many(batch)
            inserted += len(batch)
        return {"inseriti": inserted, "aggiornati": 0, "errori": errors}

    @router.post("/import/excel")
    async def import_excel(file: UploadFile = File(...)):
        try:
            import openpyxl
        except Exception:
            raise HTTPException(500, "openpyxl non installato")
        content = await read_capped(file)
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, keep_vba=False)
        inserted = updated = 0
        if "RIEP_VENDITA" in workbook.sheetnames:
            worksheet = workbook["RIEP_VENDITA"]
            for row in worksheet.iter_rows(min_row=3, values_only=True):
                code = row[0]
                if not code:
                    continue
                description = row[1] or ""
                data = {
                    "codice": str(code).strip(),
                    "descrizione": str(description).strip(),
                    "acquistati": int(row[2] or 0),
                    "venduti_negozio": int(row[3] or 0),
                    "prezzo": normalize_nonnegative_money(row[7] or 0),
                    "giacenza_vending": int(row[13] or 0),
                    "venduti_vending": int(row[17] or 0),
                }
                data["giacenza_negozio"] = physical_shop_stock_from_excel(
                    row[6], data["giacenza_vending"], data["venduti_vending"]
                )
                insert_defaults = product_model(
                    codice=data["codice"], descrizione=data["descrizione"]
                ).model_dump()
                insert_defaults = {key: value for key, value in insert_defaults.items() if key not in data}
                result = await get_database().prodotti.update_one(
                    {"codice": data["codice"]},
                    {"$set": data, "$setOnInsert": insert_defaults},
                    upsert=True,
                )
                if result.upserted_id:
                    inserted += 1
                else:
                    updated += 1
        return {"ok": True, "inseriti": inserted, "aggiornati": updated}

    @router.get("/import/excel-full/status/{job_id}")
    async def import_excel_full_status(job_id: str):
        progress = import_progress.get(job_id)
        if not progress:
            raise HTTPException(404, "Import non ancora avviato")
        return progress

    @router.post("/import/excel-full")
    async def import_excel_full(
        file: UploadFile = File(...),
        import_job_id: Optional[str] = Header(None, alias="X-Import-Job-ID"),
    ):
        try:
            import openpyxl
        except Exception:
            raise HTTPException(500, "openpyxl non installato")
        set_import_progress(import_job_id, "running", "Ricezione e controllo del file Excel", 1, 10)
        content = await read_capped(file)
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, keep_vba=False)
        except Exception as exc:
            raise HTTPException(422, f"File non leggibile: {exc}")

        expected_sheets = {
            "RIEP_VENDITA", "(SMART VENUE)", "LISTINO ADM", "RICARICA VENDING",
            "STORICO_ORDINI", "PARAMETRI", "DB_STORICO_VEND", "DB_STORICO_VENDING_EXT",
        }
        stock_sheet = None if "RIEP_VENDITA" in workbook.sheetnames else find_product_stock_sheet(workbook)
        if not expected_sheets.intersection(workbook.sheetnames) and stock_sheet is None:
            message = (
                "Nessun foglio importabile trovato. Usa il file gestionale completo oppure un foglio "
                "con le colonne CODICE, DESCRIZIONE, ACQUISTATI, GIACENZA NEGOZIO e GIACENZA VENDING."
            )
            set_import_progress(import_job_id, "failed", message, 10, 10)
            raise HTTPException(422, message)

        set_import_progress(import_job_id, "running", "Creazione del backup di sicurezza", 2, 10)
        backup = await create_backup_snapshot(
            f"Prima import {file.filename or 'Excel'}",
            "pre-import",
            excel_import_backup_collections,
        )
        found_sheets: List[str] = []
        missing_sheets: List[str] = []
        report: Dict[str, Any] = {}
        completed_steps = 0

        async def rollback_import(error: Exception) -> None:
            """Restore the explicit pre-import snapshot after a failed sheet."""
            try:
                await restore_backup_collections(backup["id"])
            except Exception:
                set_import_progress(
                    import_job_id,
                    "failed",
                    "Import fallito e rollback non riuscito: intervento necessario",
                    10,
                    10,
                    backup_id=backup["id"],
                )
                raise HTTPException(
                    500,
                    "Import fallito e rollback automatico non riuscito; ripristinare il backup "
                    f"{backup['id']}",
                ) from error
            set_import_progress(
                import_job_id,
                "failed",
                "Import fallito: dati ripristinati dal backup di sicurezza",
                10,
                10,
                backup_id=backup["id"],
            )

        async def run_sheet(name: str, importer):
            nonlocal completed_steps
            set_import_progress(
                import_job_id, "running", f"Aggiornamento foglio {name}",
                2 + completed_steps, 10, sheet=name,
            )
            if name in workbook.sheetnames:
                found_sheets.append(name)
                report[name] = await importer(workbook[name])
                errors = int(report[name].get("errori") or 0)
                if errors:
                    raise HTTPException(
                        422,
                        f"Foglio {name}: {errors} righe non valide; nessun dato è stato importato",
                    )
            else:
                missing_sheets.append(name)
            completed_steps += 1
            set_import_progress(
                import_job_id, "running", f"Foglio {name} completato",
                2 + completed_steps, 10, sheet=name, result=report.get(name),
            )

        adm_aliases = {}
        if "LISTINO ADM" in workbook.sheetnames:
            adm_aliases = adm_alias_index(
                workbook["LISTINO ADM"].iter_rows(min_row=2, values_only=True)
            )
        try:
            if "RIEP_VENDITA" in workbook.sheetnames:
                await run_sheet("RIEP_VENDITA", lambda ws: import_products(ws, adm_aliases))
                product_report_key = "RIEP_VENDITA"
            elif stock_sheet is not None:
                stock_worksheet, stock_header_row, stock_columns = stock_sheet
                product_report_key = stock_worksheet.title
                set_import_progress(
                    import_job_id, "running",
                    f"Aggiornamento giacenze dal foglio {stock_worksheet.title}",
                    2, 10, sheet=stock_worksheet.title,
                )
                found_sheets.append(stock_worksheet.title)
                report[product_report_key] = await import_product_stock_sheet(
                    stock_worksheet, stock_header_row, stock_columns, adm_aliases,
                )
                completed_steps += 1
                set_import_progress(
                    import_job_id, "running",
                    f"Giacenze dal foglio {stock_worksheet.title} completate",
                    3, 10, sheet=stock_worksheet.title, result=report[product_report_key],
                )
            else:
                await run_sheet("RIEP_VENDITA", lambda ws: import_products(ws, adm_aliases))
                product_report_key = "RIEP_VENDITA"

            await run_sheet("(SMART VENUE)", import_smart_venue)
            await run_sheet("LISTINO ADM", import_price_list)
            await run_sheet("RICARICA VENDING", import_vending)
            await run_sheet("STORICO_ORDINI", import_order_history)
            await run_sheet("PARAMETRI", import_parameters)
            await run_sheet("DB_STORICO_VEND", import_historical_sales)
            await run_sheet("DB_STORICO_VENDING_EXT", import_extended_vending_history)
        except Exception as error:
            await rollback_import(error)
            raise

        totals = {
            "prodotti_inseriti": report.get(product_report_key, {}).get("inseriti", 0),
            "prodotti_aggiornati": report.get(product_report_key, {}).get("aggiornati", 0),
            "listino_inseriti": report.get("LISTINO ADM", {}).get("inseriti", 0),
            "listino_aggiornati": report.get("LISTINO ADM", {}).get("aggiornati", 0),
            "vending_inseriti": report.get("RICARICA VENDING", {}).get("inseriti", 0),
            "vending_aggiornati": report.get("RICARICA VENDING", {}).get("aggiornati", 0),
            "smart_venue_righe": report.get("(SMART VENUE)", {}).get("inseriti", 0),
            "storico_ricreato": report.get("STORICO_ORDINI", {}).get("inseriti", 0),
            "parametri_aggiornati": report.get("PARAMETRI", {}).get("aggiornati", 0),
            "parametri_saltati": report.get("PARAMETRI", {}).get("saltati", 0),
            "db_storico_vend_righe": report.get("DB_STORICO_VEND", {}).get("inseriti", 0),
            "db_storico_vending_ext_righe": report.get("DB_STORICO_VENDING_EXT", {}).get("inseriti", 0),
            "ultimo_import_id": report.get("RIEP_VENDITA", {}).get("ultimo_import_id"),
            "valore_venduto_negozio_excel": report.get("RIEP_VENDITA", {}).get("valore_venduto_negozio_excel", 0),
            "valore_venduto_vending_excel": report.get("RIEP_VENDITA", {}).get("valore_venduto_vending_excel", 0),
            "valore_venduto_totale_excel": report.get("RIEP_VENDITA", {}).get("valore_venduto_totale_excel", 0),
        }
        response = {
            "ok": True,
            "file": file.filename,
            "backup_id": backup["id"],
            "tipo_import": "completo" if "RIEP_VENDITA" in found_sheets else "giacenze",
            "contabilita_inclusa": "RIEP_VENDITA" in found_sheets,
            "fogli_trovati": found_sheets,
            "fogli_mancanti": missing_sheets,
            "dettaglio": report,
            "totali": totals,
        }
        await record_import_history(file.filename or "Excel", response, backup["id"])
        set_import_progress(
            import_job_id, "completed", "Aggiornamento completato", 10, 10, totals=totals,
        )
        return response

    return {
        "router": router,
        "IMPORT_PROGRESS": import_progress,
        "set_import_progress": set_import_progress,
        "bulk_upsert": bulk_upsert,
        "bulk_update": bulk_update,
        "apply_electronic_inhalation_categories": apply_electronic_inhalation_categories,
        "import_products": import_products,
        "excel_header_key": excel_header_key,
        "find_product_stock_sheet": find_product_stock_sheet,
        "import_product_stock_sheet": import_product_stock_sheet,
        "import_price_list": import_price_list,
        "import_vending": import_vending,
        "import_order_history": import_order_history,
        "import_parameters": import_parameters,
        "import_historical_sales": import_historical_sales,
        "import_extended_vending_history": import_extended_vending_history,
        "import_excel": import_excel,
        "import_excel_full_status": import_excel_full_status,
        "import_excel_full": import_excel_full,
    }
