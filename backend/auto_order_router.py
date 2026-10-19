import asyncio
import hashlib
import io
import json
import logging
import math
import re
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any, Callable, Dict, List, Optional

from fastapi import APIRouter, Header, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from pymongo.errors import DuplicateKeyError


class AutoOrderRigaIn(BaseModel):
    codice: str
    quantita: int = Field(gt=0, le=100000)


class AutoOrderConfermaIn(BaseModel):
    idempotency_key: Optional[str] = None
    batch_key: Optional[str] = None
    righe: Optional[List[AutoOrderRigaIn]] = None


def auto_order_snapshot_key(ao: Dict[str, Any]) -> str:
    payload = {
        "data_riferimento_domanda": ao.get("data_riferimento_domanda"),
        "parametri": {
            key: ao.get("parametri", {}).get(key)
            for key in (
                "GIORNI_COPERTURA_MIN",
                "GIORNI_COPERTURA_TARGET",
                "AUTO_ORDER_FINESTRA_BREVE_GG",
                "AUTO_ORDER_FINESTRA_LUNGA_GG",
                "AUTO_ORDER_PESO_BREVE",
                "AUTO_ORDER_MIN_VENDUTO_BREVE",
                "AUTO_ORDER_MIN_VENDUTO_LUNGO",
                "AUTO_ORDER_FATTORE_SICUREZZA",
            )
        },
        "righe": [
            {
                "codice": row.get("codice"),
                "qta_da_ordinare": row.get("qta_da_ordinare"),
                "prezzo": row.get("prezzo"),
                "totale": row.get("totale"),
            }
            for row in ao.get("righe", [])
        ],
    }
    raw = json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str)
    return "auto-order:" + hashlib.sha256(raw.encode("utf-8")).hexdigest()


def create_auto_order_router(
    *,
    get_database: Callable[[], Any],
    get_params: Callable[[], Any],
    param_value: Callable[[Dict[str, float], str, float], float],
    lot_for: Callable[[str, Dict[str, float]], int],
    vending_stock_by_code: Callable[[], Any],
    apply_vending_stock: Callable[[List[Dict[str, Any]], Dict[str, Any]], None],
    normalize_money: Callable[[Any], float],
    sum_money: Callable[[Any], float],
    multiply_money: Callable[[Any, Any], float],
    get_auto_order: Callable[[], Any],
) -> Dict[str, Any]:
    router = APIRouter()

    @router.get("/auto-order")
    async def auto_order():
        db = get_database()
        params = await get_params()
        finestra_breve = max(3, min(30, int(param_value(params, "AUTO_ORDER_FINESTRA_BREVE_GG", 10))))
        finestra_lunga = max(
            finestra_breve,
            min(90, int(param_value(params, "AUTO_ORDER_FINESTRA_LUNGA_GG", 30))),
        )
        peso_breve = max(0.0, min(1.0, float(param_value(params, "AUTO_ORDER_PESO_BREVE", 0.70))))
        peso_lungo = 1.0 - peso_breve
        min_venduto_breve = max(0, int(param_value(params, "AUTO_ORDER_MIN_VENDUTO_BREVE", 2)))
        min_venduto_lungo = max(0, int(param_value(params, "AUTO_ORDER_MIN_VENDUTO_LUNGO", 4)))

        latest_app, latest_imported = await asyncio.gather(
            db.vendite.find_one({}, {"_id": 0, "data": 1}, sort=[("data", -1)]),
            db.db_storico_vend.find_one({}, {"_id": 0, "data": 1}, sort=[("data", -1)]),
        )

        def parsed_date(document: Optional[Dict[str, Any]]) -> Optional[datetime]:
            raw = (document or {}).get("data")
            if not raw:
                return None
            if isinstance(raw, datetime):
                return raw
            try:
                return datetime.fromisoformat(str(raw).replace("Z", "+00:00"))
            except ValueError:
                return None

        available_dates = [date for date in (parsed_date(latest_app), parsed_date(latest_imported)) if date]
        reference_date = max(available_dates, key=lambda date: date.date()) if available_dates else datetime.now(timezone.utc)
        reference_day = reference_date.date()

        def sales_pipeline(days: int) -> List[Dict[str, Any]]:
            first_day = reference_day - timedelta(days=days - 1)
            since = datetime.combine(first_day, datetime.min.time()).isoformat()
            return [
                {"$match": {"data": {"$gte": since}}},
                {"$group": {"_id": "$codice", "tot": {"$sum": "$quantita"}}},
            ]

        pipeline_breve = sales_pipeline(finestra_breve)
        pipeline_lunga = sales_pipeline(finestra_lunga)
        storico_pipeline = [
            {
                "$group": {
                    "_id": "$codice",
                    "n_ord": {"$sum": 1},
                    "tot_quantita": {"$sum": "$quantita"},
                }
            }
        ]

        (
            prodotti,
            vendite_app_breve,
            vendite_app_lunga,
            vendite_importate_breve,
            vendite_importate_lunga,
            storico_ordini,
            vending_totals,
        ) = await asyncio.gather(
            db.prodotti.find({}, {"_id": 0}).to_list(5000),
            db.vendite.aggregate(pipeline_breve).to_list(5000),
            db.vendite.aggregate(pipeline_lunga).to_list(5000),
            db.db_storico_vend.aggregate(pipeline_breve).to_list(5000),
            db.db_storico_vend.aggregate(pipeline_lunga).to_list(5000),
            db.storico_ordini.aggregate(storico_pipeline).to_list(5000),
            vending_stock_by_code(),
        )
        apply_vending_stock(prodotti, vending_totals)

        def totals(rows: List[Dict[str, Any]]) -> Dict[str, int]:
            return {row["_id"]: int(row.get("tot", 0) or 0) for row in rows if row.get("_id") is not None}

        vendite_app_breve_per_codice = totals(vendite_app_breve)
        vendite_app_lunga_per_codice = totals(vendite_app_lunga)
        vendite_importate_breve_per_codice = totals(vendite_importate_breve)
        vendite_importate_lunga_per_codice = totals(vendite_importate_lunga)
        storico_per_codice = {row["_id"]: row for row in storico_ordini if row.get("_id") is not None}

        fatt = max(1.0, float(param_value(params, "AUTO_ORDER_FATTORE_SICUREZZA", 1.15)))
        gg_min = max(1.0, float(param_value(params, "GIORNI_COPERTURA_MIN", 7)))
        gg_target = max(gg_min, float(param_value(params, "GIORNI_COPERTURA_TARGET", 14)))

        proposte = []
        esclusi = []
        for prodotto in prodotti:
            codice = prodotto["codice"]
            giac_negozio = int(prodotto.get("giacenza_negozio", 0) or 0)
            giac_vending = int(prodotto.get("giacenza_vending", 0) or 0)
            venduti_vending = int(prodotto.get("venduti_vending", 0) or 0)
            giac_tot = giac_negozio + giac_vending
            magazzino_reale_lordo = giac_negozio
            magazzino_reale = max(0, magazzino_reale_lordo)
            anomalia = giac_negozio < 0

            if codice in vendite_importate_lunga_per_codice:
                venduto_breve = vendite_importate_breve_per_codice.get(codice, 0)
                venduto_lungo = vendite_importate_lunga_per_codice.get(codice, 0)
                fonte_domanda = "DB_STORICO_VEND"
            else:
                venduto_breve = vendite_app_breve_per_codice.get(codice, 0)
                venduto_lungo = vendite_app_lunga_per_codice.get(codice, 0)
                fonte_domanda = "VENDITE_APP"

            storico = storico_per_codice.get(codice, {})
            n_ord = int(storico.get("n_ord", 0) or 0)
            tot_quantita = storico.get("tot_quantita", 0) or 0
            media_ord = round(tot_quantita / n_ord, 2) if n_ord else 0

            domanda_breve_gg = venduto_breve / float(finestra_breve)
            domanda_lunga_gg = venduto_lungo / float(finestra_lunga)
            domanda_gg = (domanda_breve_gg * peso_breve) + (domanda_lunga_gg * peso_lungo)
            lotto = lot_for(prodotto.get("categoria", "ACCESSORI"), params)
            copertura_gg = magazzino_reale / domanda_gg if domanda_gg > 0 else None
            movimento_sufficiente = venduto_breve >= min_venduto_breve or venduto_lungo >= min_venduto_lungo
            target_scorta = math.ceil(domanda_gg * gg_target * fatt) if domanda_gg > 0 else 0
            fabbisogno_grezzo = max(0, target_scorta - magazzino_reale)

            stato = "NESSUN ORDINE"
            motivo = "NESSUNA VENDITA RECENTE"
            qta = 0
            if anomalia:
                stato = "ANOMALIA"
                motivo = f"STOCK NEGOZIO FISICO NEGATIVO ({giac_negozio}) · VERIFICARE IMPORT/MOVIMENTI"
            elif not movimento_sufficiente:
                if magazzino_reale == 0 and venduto_lungo > 0:
                    stato = "CONTROLLO MANUALE"
                    motivo = (
                        f"SCORTA ZERO MA MOVIMENTO BASSO ({venduto_breve}/{venduto_lungo} "
                        f"in {finestra_breve}/{finestra_lunga}gg)"
                    )
                elif venduto_lungo > 0:
                    motivo = (
                        f"MOVIMENTO BASSO ({venduto_breve}/{venduto_lungo} "
                        f"in {finestra_breve}/{finestra_lunga}gg)"
                    )
            elif copertura_gg is not None and copertura_gg < gg_min:
                stato = "ORDINA ORA"
                motivo = f"COPERTURA {copertura_gg:.2f}gg < {int(gg_min)}gg"
                qta = max(lotto, fabbisogno_grezzo)
                qta = ((qta + lotto - 1) // lotto) * lotto
            elif copertura_gg is not None and copertura_gg < gg_target:
                stato = "MONITORA"
                motivo = f"COPERTURA {copertura_gg:.2f}gg TRA {int(gg_min)} E {int(gg_target)}gg"
            elif copertura_gg is not None:
                motivo = f"COPERTURA SUFFICIENTE ({copertura_gg:.2f}gg)"

            row = {
                "codice": codice,
                "descrizione": prodotto.get("descrizione", ""),
                "categoria": prodotto.get("categoria", ""),
                "stato": stato,
                "anomalia": anomalia,
                "qta_da_ordinare": qta,
                "giacenza": giac_tot,
                "giacenza_negozio": giac_negozio,
                "venduti_vending": venduti_vending,
                "giacenza_vending": giac_vending,
                "magazzino_reale_lordo": magazzino_reale_lordo,
                "magazzino_reale": magazzino_reale,
                "venduto_periodo": venduto_breve,
                "finestra_domanda_gg": finestra_breve,
                "venduto_breve": venduto_breve,
                "venduto_lungo": venduto_lungo,
                "venduto_10gg": venduto_breve,
                "venduto_30gg": venduto_lungo,
                "fonte_domanda": fonte_domanda,
                "domanda_gg_10": round(domanda_breve_gg, 3),
                "domanda_gg_30": round(domanda_lunga_gg, 3),
                "domanda_gg": round(domanda_gg, 3),
                "copertura_gg": round(copertura_gg, 2) if copertura_gg is not None else None,
                "target_scorta": target_scorta,
                "fabbisogno_grezzo": fabbisogno_grezzo,
                "media_ordini_storico": media_ord,
                "n_ordini_storici": n_ord,
                "lotto_ordine": lotto,
                "motivo": motivo,
                "prezzo": prodotto.get("prezzo", 0),
                "totale": multiply_money(prodotto.get("prezzo") or 0, qta),
            }
            if stato == "ORDINA ORA" and not anomalia and qta > 0:
                proposte.append(row)
            else:
                esclusi.append(row)

        proposte.sort(key=lambda row: (-row["totale"], row["codice"]))
        ordine_stati = {"ANOMALIA": 0, "CONTROLLO MANUALE": 1, "MONITORA": 2, "NESSUN ORDINE": 3}
        esclusi.sort(key=lambda row: (ordine_stati.get(row["stato"], 9), row["descrizione"], row["codice"]))
        totale = sum_money(row["totale"] for row in proposte)
        riepilogo_stati = {"ORDINA ORA": len(proposte)}
        for row in esclusi:
            riepilogo_stati[row["stato"]] = riepilogo_stati.get(row["stato"], 0) + 1
        n_anomalie_stock = sum(1 for row in [*proposte, *esclusi] if row["anomalia"])
        response = {
            "righe": proposte,
            "esclusi": esclusi,
            "totale": totale,
            "n_righe": len(proposte),
            "n_anomalie_stock": n_anomalie_stock,
            "riepilogo_stati": riepilogo_stati,
            "parametri": params,
            "finestra_domanda_gg": finestra_breve,
            "finestra_breve_gg": finestra_breve,
            "finestra_lunga_gg": finestra_lunga,
            "peso_breve": peso_breve,
            "peso_lungo": peso_lungo,
            "fattore_sicurezza": fatt,
            "data_riferimento_domanda": reference_day.isoformat(),
            "giorni_ritardo_dati": max(0, (datetime.now(timezone.utc).date() - reference_day).days),
        }
        response["snapshot_key"] = auto_order_snapshot_key(response)
        return response

    def auto_order_selected_rows(
        ao: Dict[str, Any], requested: Optional[List[AutoOrderRigaIn]] = None
    ) -> List[Dict[str, Any]]:
        if requested is None:
            return [
                dict(row)
                for row in ao["righe"]
                if not row.get("anomalia") and int(row.get("qta_da_ordinare") or 0) > 0
            ]
        if not requested:
            raise HTTPException(422, "Aggiungi almeno un articolo all'ordine")

        available = {
            row["codice"]: row
            for row in [*ao.get("righe", []), *ao.get("esclusi", [])]
            if row.get("codice") and not row.get("anomalia")
        }
        seen = set()
        selected = []
        for item in requested:
            codice = item.codice.strip()
            if codice in seen:
                raise HTTPException(422, f"Articolo duplicato nella selezione: {codice}")
            row = available.get(codice)
            if not row:
                raise HTTPException(422, f"Articolo non ordinabile: {codice}")
            seen.add(codice)
            quantity = int(item.quantita)
            selected.append(
                {
                    **row,
                    "stato": "ORDINA ORA",
                    "qta_da_ordinare": quantity,
                    "totale": multiply_money(row.get("prezzo") or 0, quantity),
                }
            )
        return selected

    def auto_order_rows_from_query(ao: Dict[str, Any], selection: Optional[str]) -> List[Dict[str, Any]]:
        if not selection:
            return auto_order_selected_rows(ao)
        try:
            raw = json.loads(selection)
            requested = [AutoOrderRigaIn.model_validate(item) for item in raw]
        except Exception as exc:
            raise HTTPException(422, "Selezione Auto-Order non valida") from exc
        return auto_order_selected_rows(ao, requested)

    @router.post("/auto-order/conferma")
    async def conferma_auto_order(
        body: Optional[AutoOrderConfermaIn] = None,
        idempotency_key: Optional[str] = Header(default=None, alias="Idempotency-Key"),
    ):
        db = get_database()
        ao = await get_auto_order()
        requested_rows = body.righe if body else None
        righe = auto_order_selected_rows(ao, requested_rows)
        selection_key = hashlib.sha256(
            json.dumps(
                [{"codice": row["codice"], "quantita": row["qta_da_ordinare"]} for row in righe],
                ensure_ascii=False,
                sort_keys=True,
            ).encode("utf-8")
        ).hexdigest()[:24]
        header_key = idempotency_key if isinstance(idempotency_key, str) else None
        default_key = (
            f"{ao['snapshot_key']}:selezione:{selection_key}"
            if requested_rows is not None
            else ao["snapshot_key"]
        )
        batch_key = (
            header_key
            or (body.idempotency_key if body else None)
            or (body.batch_key if body else None)
            or default_key
        ).strip()
        if not batch_key:
            batch_key = ao["snapshot_key"]
        existing = await db.ordini_fornitore.find_one({"batch_key": batch_key}, {"_id": 0})
        if existing:
            return {**existing, "duplicate": True}

        batch = {
            "id": str(uuid.uuid4()),
            "batch_key": batch_key,
            "snapshot_key": ao["snapshot_key"],
            "data": datetime.now(timezone.utc).isoformat(),
            "stato": "CREATO",
            "origine": "AUTO-ORDER",
            "ordinati": len(righe),
            "totale": sum_money(row.get("totale", 0) or 0 for row in righe),
            "data_riferimento_domanda": ao.get("data_riferimento_domanda"),
        }
        inserted_ids = []
        try:
            try:
                await db.ordini_fornitore.insert_one(batch.copy())
            except DuplicateKeyError:
                existing = await db.ordini_fornitore.find_one({"batch_key": batch_key}, {"_id": 0})
                if existing:
                    return {**existing, "duplicate": True}
                raise
            for index, row in enumerate(righe, start=1):
                document = {
                    "id": str(uuid.uuid4()),
                    "batch_key": batch_key,
                    "riga": index,
                    "codice": row["codice"],
                    "descrizione": row.get("descrizione", ""),
                    "categoria": row.get("categoria", ""),
                    "quantita": int(row["qta_da_ordinare"]),
                    "prezzo": normalize_money(row.get("prezzo") or 0),
                    "totale": normalize_money(row.get("totale") or 0),
                    "motivo": row.get("motivo", ""),
                }
                await db.ordini_fornitore_righe.insert_one(document)
                inserted_ids.append(document["id"])
        except Exception as exc:
            try:
                if inserted_ids:
                    await db.ordini_fornitore_righe.delete_many({"id": {"$in": inserted_ids}})
                await db.ordini_fornitore.delete_one({"batch_key": batch_key})
            except Exception:
                logging.exception("auto-order rollback failed for batch %s", batch_key)
            raise HTTPException(500, f"Creazione batch Auto-Order annullata: {exc}") from exc

        return {"ok": True, **batch, "duplicate": False}

    @router.get("/auto-order/pdf")
    async def auto_order_pdf(
        fornitore: Optional[str] = "Fornitore",
        categoria: Optional[str] = None,
        selezione: Optional[str] = None,
    ):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
        from xml.sax.saxutils import escape as xml_escape

        fornitore_safe = xml_escape((fornitore or "Fornitore")[:120])
        categoria_filtro = (categoria or "").strip()[:120]
        categoria_safe = xml_escape(categoria_filtro)

        ao = await get_auto_order()
        righe = auto_order_rows_from_query(ao, selezione)
        if categoria_filtro:
            righe = [row for row in righe if row.get("categoria") == categoria_filtro]
        totale = sum_money(row.get("totale", 0) or 0 for row in righe)

        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=15 * mm,
            rightMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm,
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "t",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=18,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=2,
        )
        subtitle_style = ParagraphStyle(
            "s",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=12,
        )
        cell_style = ParagraphStyle(
            "cell",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=7,
            leading=8,
            textColor=colors.HexColor("#0F172A"),
        )

        story = [Paragraph("ORDINE FORNITORE - GOD SERVICES", title_style)]
        today = datetime.now(timezone.utc).strftime("%d/%m/%Y")
        filter_label = f" &nbsp;·&nbsp; Selezione: <b>{categoria_safe}</b>" if categoria_filtro else ""
        story.append(
            Paragraph(
                f"Destinatario: <b>{fornitore_safe}</b> &nbsp;·&nbsp; Data: <b>{today}</b>"
                f"{filter_label} &nbsp;·&nbsp; Righe: <b>{len(righe)}</b> &nbsp;·&nbsp; "
                f"Totale: <b>€ {totale:.2f}</b>",
                subtitle_style,
            )
        )

        data = [["CODICE", "ARTICOLO", "QTA", "TOTALE"]]
        for row in righe:
            data.append(
                [
                    row["codice"],
                    Paragraph(xml_escape(row["descrizione"] or ""), cell_style),
                    str(row["qta_da_ordinare"]),
                    f"€ {row['totale']:.2f}",
                ]
            )
        data.append(["", "", "TOTALE", f"€ {totale:.2f}"])

        table = Table(data, colWidths=[25 * mm, 105 * mm, 15 * mm, 35 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 8),
                    ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                    ("ALIGN", (2, 1), (3, -1), "RIGHT"),
                    ("FONTNAME", (0, 1), (-1, -2), "Helvetica"),
                    ("FONTSIZE", (0, 1), (-1, -1), 8),
                    ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#0F172A")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#0F172A")),
                    ("LINEBELOW", (0, -2), (-1, -2), 0.5, colors.HexColor("#E2E8F0")),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F1F5F9")),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, -1), (-1, -1), 9),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)
        footer = Paragraph(
            "Calcolo Auto-Order: domanda ponderata {peso_breve:.0f}% ultimi {breve}gg + "
            "{peso_lungo:.0f}% ultimi {lungo}gg fino al {riferimento}; giacenza negozio = "
            "stock fisico libero gia normalizzato dall'import; ordine sotto {copertura}gg "
            "verso target {target}gg con sicurezza x{fattore:.2f}. Il PDF contiene solo "
            "ORDINA ORA.".format(
                peso_breve=float(ao.get("peso_breve", 0.70)) * 100,
                peso_lungo=float(ao.get("peso_lungo", 0.30)) * 100,
                breve=int(ao.get("finestra_breve_gg", 10)),
                lungo=int(ao.get("finestra_lunga_gg", 30)),
                riferimento=datetime.fromisoformat(ao["data_riferimento_domanda"]).strftime("%d/%m/%Y"),
                copertura=int(ao["parametri"].get("GIORNI_COPERTURA_MIN", 7)),
                target=int(ao["parametri"].get("GIORNI_COPERTURA_TARGET", 14)),
                fattore=float(ao.get("fattore_sicurezza", 1.15)),
            ),
            subtitle_style,
        )

        def draw_footer(canvas, current_document):
            footer.wrap(current_document.width, 14 * mm)
            footer.drawOn(canvas, current_document.leftMargin, 4 * mm)

        document.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
        buffer.seek(0)
        categoria_slug = re.sub(r"[^a-z0-9]+", "_", categoria_filtro.lower()).strip("_")
        suffix = f"_{categoria_slug}" if categoria_slug else ""
        filename = f"ordine{suffix}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M')}.pdf"
        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "X-Auto-Order-Category": categoria_filtro or "TUTTE",
                "X-Auto-Order-Rows": str(len(righe)),
            },
        )

    @router.get("/auto-order/excel")
    async def auto_order_excel(
        fornitore: Optional[str] = "Fornitore",
        categoria: Optional[str] = None,
        selezione: Optional[str] = None,
    ):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        fornitore_testo = (fornitore or "Fornitore").strip()[:120] or "Fornitore"
        categoria_filtro = (categoria or "").strip()[:120]

        ao = await get_auto_order()
        righe = auto_order_rows_from_query(ao, selezione)
        if categoria_filtro:
            righe = [row for row in righe if row.get("categoria") == categoria_filtro]
        totale = sum_money(row.get("totale", 0) or 0 for row in righe)

        def excel_text(value: Any) -> str:
            text = str(value or "")
            return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Ordine fornitore"
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A5"

        dark_fill = PatternFill("solid", fgColor="0F172A")
        light_fill = PatternFill("solid", fgColor="F1F5F9")
        white_bold = Font(color="FFFFFF", bold=True)

        worksheet.merge_cells("A1:D1")
        worksheet["A1"] = "ORDINE FORNITORE - GOD SERVICES"
        worksheet["A1"].font = Font(size=16, bold=True, color="0F172A")
        worksheet["A1"].alignment = Alignment(vertical="center")
        worksheet.row_dimensions[1].height = 26

        worksheet["A2"] = "Destinatario"
        worksheet["B2"] = excel_text(fornitore_testo)
        worksheet["C2"] = "Data"
        worksheet["D2"] = datetime.now(timezone.utc).date()
        worksheet["D2"].number_format = "dd/mm/yyyy"
        worksheet["A3"] = "Selezione"
        worksheet["B3"] = excel_text(categoria_filtro or "Tutte le categorie")
        worksheet["C3"] = "Righe"
        worksheet["D3"] = len(righe)
        for cell in (worksheet["A2"], worksheet["C2"], worksheet["A3"], worksheet["C3"]):
            cell.font = Font(bold=True, color="475569")

        header_row = 4
        for column, label in enumerate(["CODICE", "ARTICOLO", "QTA", "TOTALE"], start=1):
            cell = worksheet.cell(row=header_row, column=column, value=label)
            cell.fill = dark_fill
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="right" if column >= 3 else "left")

        first_data_row = header_row + 1
        for row_index, row in enumerate(righe, start=first_data_row):
            worksheet.cell(row=row_index, column=1, value=excel_text(row.get("codice")))
            worksheet.cell(row=row_index, column=2, value=excel_text(row.get("descrizione")))
            worksheet.cell(row=row_index, column=3, value=int(row.get("qta_da_ordinare") or 0))
            total_cell = worksheet.cell(row=row_index, column=4, value=float(row.get("totale") or 0))
            total_cell.number_format = "€ #,##0.00"
            worksheet.cell(row=row_index, column=3).alignment = Alignment(horizontal="right")
            total_cell.alignment = Alignment(horizontal="right")

        total_row = first_data_row + len(righe)
        worksheet.cell(row=total_row, column=3, value="TOTALE")
        worksheet.cell(row=total_row, column=4, value=totale)
        for column in range(1, 5):
            cell = worksheet.cell(row=total_row, column=column)
            cell.fill = light_fill
            cell.font = Font(bold=True, color="0F172A")
        worksheet.cell(row=total_row, column=4).number_format = "€ #,##0.00"
        worksheet.cell(row=total_row, column=3).alignment = Alignment(horizontal="right")
        worksheet.cell(row=total_row, column=4).alignment = Alignment(horizontal="right")

        last_filter_row = max(header_row, total_row - 1)
        worksheet.auto_filter.ref = f"A{header_row}:D{last_filter_row}"
        for column, width in {1: 20, 2: 55, 3: 14, 4: 18}.items():
            worksheet.column_dimensions[get_column_letter(column)].width = width

        workbook.properties.title = "Ordine fornitore"
        workbook.properties.subject = categoria_filtro or "Tutte le categorie"
        workbook.properties.creator = "God Services Gestionale"

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        categoria_slug = re.sub(r"[^a-z0-9]+", "_", categoria_filtro.lower()).strip("_")
        suffix = f"_{categoria_slug}" if categoria_slug else ""
        filename = f"ordine{suffix}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M')}.xlsx"
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "X-Auto-Order-Category": categoria_filtro or "TUTTE",
                "X-Auto-Order-Rows": str(len(righe)),
            },
        )

    return {
        "router": router,
        "auto_order": auto_order,
        "auto_order_selected_rows": auto_order_selected_rows,
        "auto_order_rows_from_query": auto_order_rows_from_query,
        "conferma_auto_order": conferma_auto_order,
        "auto_order_pdf": auto_order_pdf,
        "auto_order_excel": auto_order_excel,
    }
