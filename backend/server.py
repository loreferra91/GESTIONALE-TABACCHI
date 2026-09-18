from fastapi import FastAPI, APIRouter, HTTPException, UploadFile, File, Request
from fastapi.responses import FileResponse, PlainTextResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os, json, logging, uuid, io, re, base64, secrets
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Dict, Any
from datetime import datetime, timezone, timedelta

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI(title="God Services Gestionale Tabacchi")
api = APIRouter(prefix="/api")


@app.middleware("http")
async def optional_basic_auth(request: Request, call_next):
    """Protect preview deployments when APP_USERNAME/PASSWORD are configured."""
    username = os.environ.get("APP_USERNAME")
    password = os.environ.get("APP_PASSWORD")
    is_health_check = request.url.path == "/api/" or (
        request.method == "HEAD" and request.url.path == "/"
    )
    if not username or not password or is_health_check:
        return await call_next(request)

    authorization = request.headers.get("Authorization", "")
    try:
        scheme, encoded = authorization.split(" ", 1)
        decoded = base64.b64decode(encoded).decode("utf-8")
        supplied_username, supplied_password = decoded.split(":", 1)
        authenticated = (
            scheme.lower() == "basic"
            and secrets.compare_digest(supplied_username, username)
            and secrets.compare_digest(supplied_password, password)
        )
    except (ValueError, UnicodeDecodeError, base64.binascii.Error):
        authenticated = False

    if not authenticated:
        return PlainTextResponse(
            "Autenticazione richiesta",
            status_code=401,
            headers={"WWW-Authenticate": 'Basic realm="Gestionale"'},
        )
    return await call_next(request)


# ------------------------- Models -------------------------
class Prodotto(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    codice: str
    descrizione: str
    categoria: str = "ACCESSORI"  # SIGARETTE, SIGARETTE ELETTRONICHE, ACCESSORI
    prezzo: float = 0
    acquistati: int = 0
    venduti_negozio: int = 0
    venduti_vending: int = 0
    giacenza_negozio: int = 0
    giacenza_vending: int = 0
    colonna_vending: Optional[str] = None
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())


class ProdottoIn(BaseModel):
    codice: str
    descrizione: str
    categoria: str = "ACCESSORI"
    prezzo: float = 0
    acquistati: int = 0
    venduti_negozio: int = 0
    venduti_vending: int = 0
    giacenza_negozio: int = 0
    giacenza_vending: int = 0
    colonna_vending: Optional[str] = None


class ListinoItem(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    codice: str
    descrizione: str
    prezzo: float = 0
    confezione: str = ""


class VendingColonna(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    colonna: str
    codice: str = ""
    descrizione: str = ""
    giacenza: int = 0
    capacita_max: int = 5
    soglia_minima: int = 2


class VenditaGiornaliera(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    data: str
    codice: str
    descrizione: str = ""
    quantita: int = 1
    importo: float = 0
    canale: str = "NEGOZIO"  # NEGOZIO / VENDING
    pagamento: str = "CONTANTI"


class VenditaIn(BaseModel):
    data: str
    codice: str
    descrizione: Optional[str] = ""
    quantita: int = 1
    importo: float = 0
    canale: str = "NEGOZIO"
    pagamento: str = "CONTANTI"


class OrdineStorico(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    data: str
    file_sorgente: str = ""
    codice: str
    descrizione: str = ""
    quantita: int
    prezzo: float = 0


class OrdineIn(BaseModel):
    data: str
    codice: str
    descrizione: Optional[str] = ""
    quantita: int
    prezzo: float = 0
    file_sorgente: Optional[str] = "manuale"


class MovimentoCassa(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    data: str
    tipo: str  # ENTRATA / USCITA / SALDO_INIZIALE
    importo: float
    descrizione: str = ""
    operatore: str = ""


class MovimentoIn(BaseModel):
    data: str
    tipo: str
    importo: float
    descrizione: Optional[str] = ""
    operatore: Optional[str] = ""


class Parametro(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    nome: str
    valore: float
    descrizione: str = ""


class ParametroIn(BaseModel):
    valore: float


# ------------------------- Seed -------------------------
DEFAULT_PARAMS = {
    "LOTTO_SIGARETTE": {"valore": 10, "descrizione": "Lotto standard per SIGARETTE"},
    "LOTTO_ELETTRONICHE": {"valore": 5, "descrizione": "Lotto standard per SIGARETTE ELETTRONICHE"},
    "LOTTO_ACCESSORI": {"valore": 1, "descrizione": "Lotto standard per ACCESSORI"},
    "GIORNI_STORICO_VEND": {"valore": 30, "descrizione": "Finestra storico vendite (giorni)"},
    "GIORNI_SETTIMANA": {"valore": 7, "descrizione": "Costante giorni settimana"},
    "VENDITE_GIORNALIERE_MESE": {"valore": 6.5, "descrizione": "Divisore vendite medie giornaliere/mese"},
    "SOGLIA_ALLERT_PCT": {"valore": 0.35, "descrizione": "Soglia % giacenza per allert riordino"},
    "FAST_VENDUTO30_MIN": {"valore": 8, "descrizione": "Venduto min 30gg per FAST MOVER"},
    "SLOW_VENDUTO30_MAX": {"valore": 2, "descrizione": "Venduto max 30gg per SLOW MOVER"},
    "SLOW_TARGET_FACTOR": {"valore": 0.6, "descrizione": "Fattore target settimanale slow mover"},
    "FATT_SETTIMANALE": {"valore": 1.15, "descrizione": "Fattore fabbisogno settimanale"},
}


async def seed_if_empty():
    seed_path = ROOT_DIR / "seed" / "seed_data.json"
    if not seed_path.exists():
        return
    count = await db.prodotti.count_documents({})
    if count > 0:
        return
    with open(seed_path, encoding="utf-8") as f:
        data = json.load(f)
    if data.get("prodotti"):
        docs = [Prodotto(**p).model_dump() for p in data["prodotti"]]
        await db.prodotti.insert_many(docs)
    if data.get("listino_adm"):
        docs = [ListinoItem(**p).model_dump() for p in data["listino_adm"]]
        # chunked insert
        for i in range(0, len(docs), 1000):
            await db.listino_adm.insert_many(docs[i:i+1000])
    if data.get("vending"):
        docs = [VendingColonna(**p).model_dump() for p in data["vending"]]
        await db.vending.insert_many(docs)
    if data.get("storico_ordini"):
        docs = [OrdineStorico(**p).model_dump() for p in data["storico_ordini"]]
        for i in range(0, len(docs), 1000):
            await db.storico_ordini.insert_many(docs[i:i+1000])
    # Parametri
    seed_params = {**DEFAULT_PARAMS, **{k: {"valore": v["valore"] if isinstance(v, dict) else v, "descrizione": v.get("descrizione","") if isinstance(v, dict) else ""} for k, v in (data.get("parametri") or {}).items()}}
    param_docs = [Parametro(nome=k, valore=float(v["valore"]) if v["valore"] is not None else 0, descrizione=v.get("descrizione","")).model_dump() for k, v in seed_params.items()]
    await db.parametri.insert_many(param_docs)


@app.on_event("startup")
async def on_start():
    try:
        await seed_if_empty()
        # ensure parametri exists if empty
        if await db.parametri.count_documents({}) == 0:
            docs = [Parametro(nome=k, valore=float(v["valore"]), descrizione=v["descrizione"]).model_dump() for k, v in DEFAULT_PARAMS.items()]
            await db.parametri.insert_many(docs)
    except Exception as e:
        logging.exception("seed failed: %s", e)


# ------------------------- Helpers -------------------------
def strip_id(d):
    if d and "_id" in d:
        d.pop("_id", None)
    return d


async def get_params() -> Dict[str, float]:
    docs = await db.parametri.find({}, {"_id": 0}).to_list(1000)
    return {d["nome"]: float(d["valore"]) for d in docs}


async def compute_venduto_30gg(codice: str) -> int:
    since = (datetime.now(timezone.utc) - timedelta(days=30)).isoformat()
    pipeline = [
        {"$match": {"codice": codice, "data": {"$gte": since}}},
        {"$group": {"_id": None, "tot": {"$sum": "$quantita"}}},
    ]
    res = await db.vendite.aggregate(pipeline).to_list(1)
    return int(res[0]["tot"]) if res else 0


def lotto_for(categoria: str, params: Dict[str, float]) -> int:
    if categoria == "SIGARETTE":
        return int(params.get("LOTTO_SIGARETTE", 10))
    if categoria == "SIGARETTE ELETTRONICHE":
        return int(params.get("LOTTO_ELETTRONICHE", 5))
    return int(params.get("LOTTO_ACCESSORI", 1))


# ------------------------- Root / Health -------------------------
@api.get("/")
async def root():
    return {"app": "God Services Gestionale Tabacchi", "status": "ok"}


# ------------------------- Prodotti -------------------------
MAX_LIMIT = 5000
MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20 MB


def _q_regex(q: str) -> Dict[str, Any]:
    """Safe regex from user input (escaped, bounded)."""
    return {"$regex": re.escape(q[:200]), "$options": "i"}


def _cap(limit: int) -> int:
    return max(1, min(int(limit or 0), MAX_LIMIT))


@api.get("/prodotti")
async def list_prodotti(q: Optional[str] = None, categoria: Optional[str] = None, limit: int = 500):
    filt: Dict[str, Any] = {}
    if categoria:
        filt["categoria"] = categoria
    if q:
        rx = _q_regex(q)
        filt["$or"] = [{"codice": rx}, {"descrizione": rx}]
    docs = await db.prodotti.find(filt, {"_id": 0}).sort("codice", 1).to_list(_cap(limit))
    return docs


@api.post("/prodotti")
async def create_prodotto(p: ProdottoIn):
    prod = Prodotto(**p.model_dump())
    await db.prodotti.insert_one(prod.model_dump())
    return prod.model_dump()


@api.put("/prodotti/{prod_id}")
async def update_prodotto(prod_id: str, p: ProdottoIn):
    r = await db.prodotti.update_one({"id": prod_id}, {"$set": p.model_dump()})
    if r.matched_count == 0:
        raise HTTPException(404, "not found")
    doc = await db.prodotti.find_one({"id": prod_id}, {"_id": 0})
    return doc


@api.delete("/prodotti/{prod_id}")
async def delete_prodotto(prod_id: str):
    await db.prodotti.delete_one({"id": prod_id})
    return {"ok": True}


# ------------------------- Listino ADM -------------------------
@api.get("/listino")
async def list_listino(q: Optional[str] = None, limit: int = 300):
    filt: Dict[str, Any] = {}
    if q:
        rx = _q_regex(q)
        filt["$or"] = [{"codice": rx}, {"descrizione": rx}]
    docs = await db.listino_adm.find(filt, {"_id": 0}).sort("descrizione", 1).to_list(_cap(limit))
    total = await db.listino_adm.count_documents(filt)
    return {"items": docs, "total": total}


# ------------------------- Vendite giornaliere -------------------------
@api.get("/vendite")
async def list_vendite(giorno: Optional[str] = None, limit: int = 500):
    filt: Dict[str, Any] = {}
    if giorno:
        # data ISO stringa (2026-02-18...) — validiamo formato semplice
        safe = re.sub(r"[^0-9\-T:.]", "", giorno)[:32]
        filt["data"] = {"$regex": f"^{re.escape(safe)}"}
    docs = await db.vendite.find(filt, {"_id": 0}).sort("data", -1).to_list(_cap(limit))
    return docs


@api.post("/vendite")
async def add_vendita(v: VenditaIn):
    vv = VenditaGiornaliera(**v.model_dump())
    # aggiorna giacenze del prodotto
    prod = await db.prodotti.find_one({"codice": v.codice})
    if prod:
        vv.descrizione = vv.descrizione or prod.get("descrizione", "")
        field = "giacenza_vending" if v.canale == "VENDING" else "giacenza_negozio"
        vend_field = "venduti_vending" if v.canale == "VENDING" else "venduti_negozio"
        await db.prodotti.update_one({"codice": v.codice}, {"$inc": {field: -v.quantita, vend_field: v.quantita}})
    await db.vendite.insert_one(vv.model_dump())
    return vv.model_dump()


@api.delete("/vendite/{v_id}")
async def del_vendita(v_id: str):
    v = await db.vendite.find_one({"id": v_id})
    if v:
        field = "giacenza_vending" if v.get("canale") == "VENDING" else "giacenza_negozio"
        vend_field = "venduti_vending" if v.get("canale") == "VENDING" else "venduti_negozio"
        await db.prodotti.update_one({"codice": v["codice"]}, {"$inc": {field: v["quantita"], vend_field: -v["quantita"]}})
    await db.vendite.delete_one({"id": v_id})
    return {"ok": True}


class BulkVenditaIn(BaseModel):
    canale: str = "NEGOZIO"
    pagamento: str = "CONTANTI"
    righe: List[Dict[str, Any]]  # {data, codice, descrizione, quantita, importo}


@api.post("/vendite/bulk")
async def bulk_vendite(body: BulkVenditaIn):
    """Bulk paste da Excel: rows with data, codice, descrizione, quantita, importo."""
    inserted = 0
    skipped = 0
    errors = []
    for i, r in enumerate(body.righe):
        try:
            codice = str(r.get("codice") or "").strip()
            if not codice:
                skipped += 1
                continue
            qta = int(float(r.get("quantita") or 0))
            imp = float(r.get("importo") or 0)
            if qta <= 0:
                skipped += 1
                continue
            data = str(r.get("data") or datetime.now(timezone.utc).isoformat())
            desc = str(r.get("descrizione") or "")
            prod = await db.prodotti.find_one({"codice": codice})
            if prod:
                desc = desc or prod.get("descrizione", "")
                field = "giacenza_vending" if body.canale == "VENDING" else "giacenza_negozio"
                vend_field = "venduti_vending" if body.canale == "VENDING" else "venduti_negozio"
                await db.prodotti.update_one({"codice": codice}, {"$inc": {field: -qta, vend_field: qta}})
            v = VenditaGiornaliera(data=data, codice=codice, descrizione=desc, quantita=qta, importo=imp, canale=body.canale, pagamento=body.pagamento)
            await db.vendite.insert_one(v.model_dump())
            inserted += 1
        except Exception as e:
            errors.append({"riga": i + 1, "errore": str(e)})
    return {"inseriti": inserted, "saltati": skipped, "errori": errors}


async def _read_capped(file: UploadFile, max_bytes: int = MAX_UPLOAD_BYTES) -> bytes:
    """Read an UploadFile with a hard size cap; raise 413 if exceeded."""
    buf = bytearray()
    while True:
        chunk = await file.read(1024 * 64)
        if not chunk:
            break
        buf.extend(chunk)
        if len(buf) > max_bytes:
            raise HTTPException(413, f"File troppo grande (max {max_bytes // 1024 // 1024} MB)")
    return bytes(buf)


@api.post("/vendite/import-csv-vending")
async def import_csv_vending(file: UploadFile = File(...), pagamento: str = "CONTANTI"):
    """Import CSV del distributore vending. Colonne accettate (case-insensitive):
       data, prodotto/nome prodotto, prezzo, colonna, codice/codice aams, categoria, pagamento.
       Separatore auto-rilevato (, ; \\t).
    """
    import csv
    raw = (await _read_capped(file)).decode("utf-8-sig", errors="replace")
    # rileva delimitatore
    sample = raw[:2000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except Exception:
        delim = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(raw), delimiter=delim)
    # normalizza header
    def norm(s): return re.sub(r"\s+", " ", (s or "").strip().lower())
    field_map = {norm(k): k for k in (reader.fieldnames or [])}
    def pick(row, *keys):
        for k in keys:
            actual = field_map.get(norm(k))
            if actual and row.get(actual) not in (None, ""):
                return row.get(actual)
        return None

    inserted = 0
    skipped = 0
    errors = []
    for i, row in enumerate(reader):
        try:
            codice = str(pick(row, "codice", "codice aams", "cod aams", "cod", "aams") or "").strip()
            nome = str(pick(row, "nome prodotto", "prodotto", "descrizione", "articolo") or "").strip()
            prezzo = pick(row, "prezzo", "importo")
            data = pick(row, "data", "date")
            colonna = str(pick(row, "colonna", "column") or "").strip()
            categoria = str(pick(row, "categoria", "tipo") or "").strip()
            pag = str(pick(row, "pagamento", "payment") or pagamento).strip() or pagamento

            if not codice and not nome:
                skipped += 1
                continue
            prezzo_f = float(str(prezzo).replace(",", ".")) if prezzo not in (None, "") else 0.0
            # data: default oggi se mancante
            data_iso = str(data) if data else datetime.now(timezone.utc).isoformat()
            # se codice manca prova a risolvere dal nome
            if not codice and nome:
                p = await db.prodotti.find_one({"descrizione": {"$regex": f"^{re.escape(nome)}$", "$options": "i"}})
                if p:
                    codice = p["codice"]
            if not codice:
                # crea un placeholder
                codice = f"CSV-{norm(nome)[:20].replace(' ', '_')}"

            # ogni riga CSV = 1 pezzo venduto (formato tipico distributore)
            v = VenditaGiornaliera(
                data=data_iso, codice=codice, descrizione=nome, quantita=1, importo=prezzo_f,
                canale="VENDING", pagamento=pag,
            )
            await db.vendite.insert_one(v.model_dump())

            # aggiorna vending column giacenza se colonna presente
            if colonna:
                col = await db.vending.find_one({"colonna": colonna})
                if col:
                    new_g = max(0, (col.get("giacenza") or 0) - 1)
                    await db.vending.update_one({"colonna": colonna}, {"$set": {"giacenza": new_g}})
            # aggiorna prodotto
            prod = await db.prodotti.find_one({"codice": codice})
            if prod:
                await db.prodotti.update_one({"codice": codice}, {"$inc": {"giacenza_vending": -1, "venduti_vending": 1}})
            else:
                # crea prodotto minimale
                await db.prodotti.insert_one(Prodotto(
                    codice=codice, descrizione=nome or codice,
                    categoria=(categoria.upper() or "SIGARETTE") if categoria else "SIGARETTE",
                    prezzo=prezzo_f, venduti_vending=1,
                ).model_dump())
            inserted += 1
        except Exception as e:
            errors.append({"riga": i + 2, "errore": str(e)})
    return {"inseriti": inserted, "saltati": skipped, "errori": errors, "delimitatore": delim}




# ------------------------- Vending -------------------------
@api.get("/vending")
async def list_vending():
    docs = await db.vending.find({}, {"_id": 0}).sort("colonna", 1).to_list(500)
    # arricchisci con esito/proposta
    out = []
    for d in docs:
        cap = d.get("capacita_max", 5) or 5
        giac = d.get("giacenza", 0) or 0
        soglia = d.get("soglia_minima", 2) or 2
        proposta = max(0, cap - giac) if giac < soglia else 0
        # esito
        if giac >= cap:
            esito = "PIENO" if giac == cap else "OLTRE CAPACITA"
        elif giac < soglia:
            esito = "DA CARICARE"
        else:
            esito = "OK"
        d["proposta"] = proposta
        d["esito"] = esito
        out.append(d)
    return out


@api.put("/vending/{v_id}")
async def update_vending(v_id: str, body: Dict[str, Any]):
    allowed = {k: body[k] for k in ("codice", "descrizione", "giacenza", "capacita_max", "soglia_minima") if k in body}
    r = await db.vending.update_one({"id": v_id}, {"$set": allowed})
    if r.matched_count == 0:
        raise HTTPException(404, "not found")
    doc = await db.vending.find_one({"id": v_id}, {"_id": 0})
    return doc


@api.post("/vending/{v_id}/ricarica")
async def ricarica_vending(v_id: str, body: Dict[str, Any]):
    qta = int(body.get("quantita", 0))
    v = await db.vending.find_one({"id": v_id})
    if not v:
        raise HTTPException(404, "not found")
    nuovo = min(v["capacita_max"], v["giacenza"] + qta)
    await db.vending.update_one({"id": v_id}, {"$set": {"giacenza": nuovo}})
    # scala dal magazzino negozio
    if v.get("codice"):
        await db.prodotti.update_one({"codice": v["codice"]}, {"$inc": {"giacenza_negozio": -qta, "giacenza_vending": qta}})
    return {"ok": True, "colonna": v["colonna"], "nuova_giacenza": nuovo}


# ------------------------- Storico ordini -------------------------
@api.get("/ordini")
async def list_ordini(limit: int = 1000):
    docs = await db.storico_ordini.find({}, {"_id": 0}).sort("data", -1).to_list(_cap(limit))
    return docs


@api.post("/ordini")
async def add_ordine(o: OrdineIn):
    oo = OrdineStorico(**o.model_dump())
    await db.storico_ordini.insert_one(oo.model_dump())
    # aggiorna giacenza / acquisti
    if o.codice:
        await db.prodotti.update_one({"codice": o.codice}, {"$inc": {"giacenza_negozio": o.quantita, "acquistati": o.quantita}})
    return oo.model_dump()


@api.delete("/ordini/{o_id}")
async def del_ordine(o_id: str):
    await db.storico_ordini.delete_one({"id": o_id})
    return {"ok": True}


class BulkOrdineIn(BaseModel):
    file_sorgente: str = "carico manuale"
    data: Optional[str] = None
    righe: List[Dict[str, Any]]  # {codice, descrizione, quantita, prezzo}


@api.post("/ordini/bulk")
async def bulk_carico(body: BulkOrdineIn):
    """Carico merce: registra multiple righe d'ordine e aggiorna giacenze."""
    data = body.data or datetime.now(timezone.utc).isoformat()
    inserted = 0
    creati_prodotti = 0
    errors = []
    for i, r in enumerate(body.righe):
        try:
            codice = str(r.get("codice") or "").strip()
            qta = int(float(r.get("quantita") or 0))
            if not codice or qta == 0:
                continue
            prezzo = float(r.get("prezzo") or 0)
            desc = str(r.get("descrizione") or "")
            o = OrdineStorico(data=data, file_sorgente=body.file_sorgente, codice=codice, descrizione=desc, quantita=qta, prezzo=prezzo)
            await db.storico_ordini.insert_one(o.model_dump())
            prod = await db.prodotti.find_one({"codice": codice})
            if prod:
                await db.prodotti.update_one(
                    {"codice": codice},
                    {"$inc": {"giacenza_negozio": qta, "acquistati": qta}},
                )
            else:
                # crea prodotto minimale
                await db.prodotti.insert_one(Prodotto(
                    codice=codice, descrizione=desc or codice,
                    categoria="SIGARETTE" if codice.startswith("AMMS") else "ACCESSORI",
                    prezzo=prezzo, acquistati=qta, giacenza_negozio=qta,
                ).model_dump())
                creati_prodotti += 1
            inserted += 1
        except Exception as e:
            errors.append({"riga": i + 1, "errore": str(e)})
    return {"caricate": inserted, "prodotti_nuovi": creati_prodotti, "errori": errors}




# ------------------------- Cassa -------------------------
@api.get("/cassa")
async def list_cassa(limit: int = 500):
    docs = await db.cassa.find({}, {"_id": 0}).sort("data", -1).to_list(_cap(limit))
    saldo = 0.0
    all_docs = await db.cassa.find({}, {"_id": 0}).to_list(10000)
    for d in all_docs:
        if d["tipo"] in ("ENTRATA", "SALDO_INIZIALE"):
            saldo += d["importo"]
        else:
            saldo -= d["importo"]
    return {"movimenti": docs, "saldo": round(saldo, 2)}


@api.post("/cassa")
async def add_cassa(m: MovimentoIn):
    mm = MovimentoCassa(**m.model_dump())
    await db.cassa.insert_one(mm.model_dump())
    return mm.model_dump()


@api.delete("/cassa/{m_id}")
async def del_cassa(m_id: str):
    await db.cassa.delete_one({"id": m_id})
    return {"ok": True}


# ------------------------- Parametri -------------------------
@api.get("/parametri")
async def list_parametri():
    docs = await db.parametri.find({}, {"_id": 0}).sort("nome", 1).to_list(200)
    return docs


@api.put("/parametri/{nome}")
async def update_parametro(nome: str, body: ParametroIn):
    r = await db.parametri.update_one({"nome": nome}, {"$set": {"valore": body.valore}})
    if r.matched_count == 0:
        # create if missing
        await db.parametri.insert_one(Parametro(nome=nome, valore=body.valore).model_dump())
    doc = await db.parametri.find_one({"nome": nome}, {"_id": 0})
    return doc


# ------------------------- Auto-Order -------------------------
@api.get("/auto-order")
async def auto_order():
    params = await get_params()
    soglia_pct = params.get("SOGLIA_ALLERT_PCT", 0.35)
    fast_min = params.get("FAST_VENDUTO30_MIN", 8)
    slow_max = params.get("SLOW_VENDUTO30_MAX", 2)
    fatt = params.get("FATT_SETTIMANALE", 1.15)

    prodotti = await db.prodotti.find({}, {"_id": 0}).to_list(5000)
    proposte = []
    for p in prodotti:
        codice = p["codice"]
        giac_tot = (p.get("giacenza_negozio", 0) or 0) + (p.get("giacenza_vending", 0) or 0)
        acq = p.get("acquistati", 0) or 0
        vend_30gg = await compute_venduto_30gg(codice)
        # storico ordini per prodotto
        storici = await db.storico_ordini.find({"codice": codice}, {"_id": 0}).to_list(1000)
        n_ord = len(storici)
        media_ord = round(sum(o.get("quantita", 0) for o in storici) / n_ord, 2) if n_ord else 0
        lotto = lotto_for(p.get("categoria", "ACCESSORI"), params)

        pct_giac = (giac_tot / acq) if acq > 0 else (1 if giac_tot > 0 else 0)

        # fabbisogno grezzo: se sotto soglia
        motivo = None
        qta = 0
        if giac_tot <= 0:
            if vend_30gg >= fast_min:
                motivo = "FAST MOVER: REINTEGRO"
                qta = max(lotto, int(round(vend_30gg * fatt / 7 * 7)))
            elif vend_30gg <= slow_max:
                motivo = "SLOW MOVER: 1 LOTTO (STOCK ZERO)"
                qta = lotto
            else:
                motivo = "TARGET SETTIMANALE"
                qta = lotto
        elif pct_giac < soglia_pct and vend_30gg > 0:
            if vend_30gg >= fast_min:
                motivo = "FAST MOVER: REINTEGRO"
                qta = max(lotto, int(round(vend_30gg * fatt)))
            else:
                motivo = "TARGET SETTIMANALE"
                qta = lotto

        if qta > 0:
            # arrotonda al lotto
            qta = ((qta + lotto - 1) // lotto) * lotto
            proposte.append({
                "codice": codice,
                "descrizione": p.get("descrizione", ""),
                "categoria": p.get("categoria", ""),
                "qta_da_ordinare": qta,
                "giacenza": giac_tot,
                "venduto_30gg": vend_30gg,
                "media_ordini_storico": media_ord,
                "n_ordini_storici": n_ord,
                "lotto_ordine": lotto,
                "motivo": motivo,
                "prezzo": p.get("prezzo", 0),
                "totale": round(qta * (p.get("prezzo") or 0), 2),
            })

    proposte.sort(key=lambda x: (-x["totale"], x["codice"]))
    tot = round(sum(x["totale"] for x in proposte), 2)
    return {"righe": proposte, "totale": tot, "n_righe": len(proposte), "parametri": params}


@api.post("/auto-order/conferma")
async def conferma_auto_order():
    """Trasforma le proposte in ordini reali (storico_ordini) e aggiorna scorte."""
    ao = await auto_order()
    oggi = datetime.now(timezone.utc).isoformat()
    inserted = 0
    for r in ao["righe"]:
        o = OrdineStorico(
            data=oggi,
            file_sorgente="AUTO-ORDER",
            codice=r["codice"],
            descrizione=r["descrizione"],
            quantita=r["qta_da_ordinare"],
            prezzo=r["prezzo"],
        )
        await db.storico_ordini.insert_one(o.model_dump())
        await db.prodotti.update_one(
            {"codice": r["codice"]},
            {"$inc": {"giacenza_negozio": r["qta_da_ordinare"], "acquistati": r["qta_da_ordinare"]}},
        )
        inserted += 1
    return {"ok": True, "ordinati": inserted, "totale": ao["totale"]}


# ------------------------- Auto-Order PDF -------------------------
@api.get("/auto-order/pdf")
async def auto_order_pdf(fornitore: Optional[str] = "Fornitore"):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.enums import TA_RIGHT, TA_LEFT
    from xml.sax.saxutils import escape as xml_escape

    # sanifica input utente per la markup di reportlab
    fornitore_safe = xml_escape((fornitore or "Fornitore")[:120])

    ao = await auto_order()
    righe = ao["righe"]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    title_s = ParagraphStyle('t', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, textColor=colors.HexColor('#0F172A'), spaceAfter=2)
    sub_s = ParagraphStyle('s', parent=styles['Normal'], fontName='Helvetica', fontSize=9, textColor=colors.HexColor('#64748B'), spaceAfter=12)
    right_s = ParagraphStyle('r', parent=styles['Normal'], fontName='Helvetica', fontSize=9, alignment=TA_RIGHT, textColor=colors.HexColor('#334155'))

    story = []
    now = datetime.now(timezone.utc).strftime("%d/%m/%Y")
    story.append(Paragraph("ORDINE FORNITORE — GOD SERVICES", title_s))
    story.append(Paragraph(f"Destinatario: <b>{fornitore_safe}</b> &nbsp;·&nbsp; Data: <b>{now}</b> &nbsp;·&nbsp; Righe: <b>{len(righe)}</b> &nbsp;·&nbsp; Totale: <b>€ {ao['totale']:.2f}</b>", sub_s))

    # Table
    header = ["CODICE", "ARTICOLO", "TIPO", "QTA", "LOTTO", "PREZZO", "TOTALE", "MOTIVO"]
    data = [header]
    for r in righe:
        data.append([
            r["codice"],
            (r["descrizione"] or "")[:45],
            (r["categoria"] or "")[:3],
            str(r["qta_da_ordinare"]),
            str(r["lotto_ordine"]),
            f"€ {r['prezzo']:.2f}",
            f"€ {r['totale']:.2f}",
            (r["motivo"] or "")[:22],
        ])
    data.append(["", "", "", "", "", "TOTALE", f"€ {ao['totale']:.2f}", ""])

    col_widths = [22*mm, 60*mm, 12*mm, 12*mm, 12*mm, 18*mm, 20*mm, 30*mm]
    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0F172A')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 8),
        ('ALIGN', (0,0), (-1,0), 'LEFT'),
        ('ALIGN', (3,1), (6,-1), 'RIGHT'),
        ('FONTNAME', (0,1), (-1,-2), 'Helvetica'),
        ('FONTSIZE', (0,1), (-1,-1), 8),
        ('TEXTCOLOR', (0,1), (-1,-1), colors.HexColor('#0F172A')),
        ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#F8FAFC')]),
        ('LINEBELOW', (0,0), (-1,0), 0.8, colors.HexColor('#0F172A')),
        ('LINEBELOW', (0,-2), (-1,-2), 0.5, colors.HexColor('#E2E8F0')),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#F1F5F9')),
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,-1), (-1,-1), 9),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 20))
    story.append(Paragraph("Documento generato automaticamente dal Gestionale God Services · Parametri: SOGLIA {sp}%, finestra {gg}gg".format(
        sp=int((ao['parametri'].get('SOGLIA_ALLERT_PCT', 0.35))*100),
        gg=int(ao['parametri'].get('GIORNI_STORICO_VEND', 30))
    ), sub_s))

    doc.build(story)
    buf.seek(0)
    fname = f"ordine_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={fname}"})


# ------------------------- Prodotti più venduti (per POS) -------------------------
@api.get("/prodotti/top")
async def prodotti_top(limit: int = 40):
    prods = await db.prodotti.find({}, {"_id": 0}).to_list(5000)
    for p in prods:
        p["venduti_totale"] = (p.get("venduti_negozio") or 0) + (p.get("venduti_vending") or 0)
    prods.sort(key=lambda x: -x["venduti_totale"])
    return prods[:limit]



# ------------------------- Pivot magazzino -------------------------
@api.get("/pivot")
async def pivot():
    prods = await db.prodotti.find({}, {"_id": 0}).to_list(5000)
    params = await get_params()
    divisor = params.get("VENDITE_GIORNALIERE_MESE", 6.5)
    rows = []
    tot_acq = tot_vend = tot_giac = pezzi_mag = 0
    fermi = negativi = da_riord = lenti = 0
    for p in prods:
        v_neg = p.get("venduti_negozio", 0) or 0
        v_ven = p.get("venduti_vending", 0) or 0
        v_tot = v_neg + v_ven
        g_neg = p.get("giacenza_negozio", 0) or 0
        g_ven = p.get("giacenza_vending", 0) or 0
        g_tot = g_neg + g_ven
        prz = p.get("prezzo", 0) or 0
        acq = p.get("acquistati", 0) or 0
        ta = acq * prz
        tv = v_tot * prz
        tg = g_tot * prz
        vend_mese = v_tot / max(1, divisor) if divisor else 0
        mesi_smalt = round(g_tot / vend_mese, 2) if vend_mese > 0 else 0
        stato = "OK"
        azione = "OK"
        if g_tot <= 0:
            stato = "ESAURITO"
            azione = "RIORDINO"
        elif v_tot == 0 and acq > 0:
            stato = "FERMO"
            azione = "REVISIONE"
            fermi += 1
        elif mesi_smalt > 6:
            stato = "LENTO"
            azione = "REVISIONE"
            lenti += 1
        if g_tot < 0:
            negativi += 1
        if stato == "ESAURITO":
            da_riord += 1
        tot_acq += ta
        tot_vend += tv
        tot_giac += tg
        pezzi_mag += g_tot
        rows.append({
            "codice": p["codice"], "descrizione": p.get("descrizione", ""),
            "categoria": p.get("categoria", ""),
            "acq": acq, "vend_negozio": v_neg, "vend_vending": v_ven, "vend_totale": v_tot,
            "giac_negozio": g_neg, "giac_vending": g_ven, "giac_totale": g_tot,
            "prezzo": prz, "tot_acquistato": round(ta, 2), "tot_venduto": round(tv, 2), "tot_giacenza": round(tg, 2),
            "vendite_mese": round(vend_mese, 2), "mesi_smaltimento": mesi_smalt,
            "stato": stato, "azione": azione,
        })
    rows.sort(key=lambda x: -x["tot_giacenza"])
    return {
        "kpi": {
            "valore_acquistato": round(tot_acq * 0.90, 2),
            "valore_venduto": round(tot_vend, 2),
            "valore_giacenza": round(tot_giac, 2),
            "pezzi_magazzino": pezzi_mag,
            "prodotti_totali": len(rows),
            "prodotti_fermi": fermi,
            "giacenza_negativa": negativi,
            "da_riordinare": da_riord,
            "lenti_oltre_6_mesi": lenti,
        },
        "righe": rows,
    }


# ------------------------- Dashboard KPIs -------------------------
@api.get("/dashboard")
async def dashboard():
    pv = await pivot()
    # vendite oggi
    oggi = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    vendite_oggi = await db.vendite.aggregate([
        {"$match": {"data": {"$regex": f"^{oggi}"}}},
        {"$group": {"_id": None, "tot": {"$sum": "$importo"}, "pezzi": {"$sum": "$quantita"}}}
    ]).to_list(1)
    v_oggi = vendite_oggi[0] if vendite_oggi else {"tot": 0, "pezzi": 0}
    # cassa
    cassa = await list_cassa()
    # vending
    vending = await list_vending()
    vend_da_caricare = sum(1 for v in vending if v["esito"] == "DA CARICARE")
    return {
        "kpi": pv["kpi"],
        "vendite_oggi": {"importo": round(v_oggi.get("tot") or 0, 2), "pezzi": v_oggi.get("pezzi") or 0},
        "saldo_cassa": cassa["saldo"],
        "vending_da_caricare": vend_da_caricare,
        "vending_totale": len(vending),
    }


# ------------------------- Import Excel -------------------------
@api.post("/import/excel")
async def import_excel(file: UploadFile = File(...)):
    try:
        import openpyxl
    except Exception:
        raise HTTPException(500, "openpyxl non installato")
    content = await _read_capped(file)
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, keep_vba=False)
    inserted = updated = 0
    if "RIEP_VENDITA" in wb.sheetnames:
        ws = wb["RIEP_VENDITA"]
        for row in ws.iter_rows(min_row=3, values_only=True):
            codice = row[0]
            if not codice:
                continue
            desc = row[1] or ""
            data_p = {
                "codice": str(codice).strip(),
                "descrizione": str(desc).strip(),
                "acquistati": int(row[2] or 0),
                "venduti_negozio": int(row[3] or 0),
                "giacenza_negozio": int(row[6] or 0),
                "prezzo": float(row[7] or 0),
                "giacenza_vending": int(row[13] or 0),
                "venduti_vending": int(row[17] or 0),
            }
            r = await db.prodotti.update_one({"codice": str(codice).strip()}, {"$set": data_p}, upsert=True)
            if r.upserted_id:
                inserted += 1
            else:
                updated += 1
    return {"ok": True, "inseriti": inserted, "aggiornati": updated}


# ------------------------- Register -------------------------
app.include_router(api)

_cors_origins = [o.strip() for o in os.environ.get('CORS_ORIGINS', '*').split(',') if o.strip()]
_cors_credentials = _cors_origins != ['*']  # wildcard + credentials è invalido; disabilita credentials se wildcard

app.add_middleware(
    CORSMiddleware,
    allow_credentials=_cors_credentials,
    allow_origins=_cors_origins,
    allow_methods=["*"],
    allow_headers=["*"],
)


# In production the React build is copied alongside the backend. API routes
# keep precedence; every other unknown path falls back to the SPA entry point.
FRONTEND_BUILD_DIR = ROOT_DIR.parent / "frontend" / "build"
if FRONTEND_BUILD_DIR.exists():
    static_dir = FRONTEND_BUILD_DIR / "static"
    if static_dir.exists():
        app.mount("/static", StaticFiles(directory=static_dir), name="static")

    @app.get("/{full_path:path}", include_in_schema=False)
    async def serve_frontend(full_path: str):
        build_root = FRONTEND_BUILD_DIR.resolve()
        requested = (build_root / full_path).resolve()
        try:
            requested.relative_to(build_root)
        except ValueError:
            requested = build_root / "index.html"

        if requested.is_file():
            return FileResponse(requested)
        return FileResponse(build_root / "index.html")

logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s %(name)s - %(message)s')


@app.on_event("shutdown")
async def _shutdown():
    client.close()
