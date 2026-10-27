"""Manual merchandise receipt history and Smart Venue confirmation routes."""
from datetime import datetime, timezone
import hashlib
import io
import json
import uuid
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, File, HTTPException, Request, UploadFile
from pydantic import BaseModel, ConfigDict, Field
from pymongo import UpdateOne
from pymongo.errors import DuplicateKeyError

try:
    from .money import NonNegativeMoney
except ImportError:  # pragma: no cover - direct module execution compatibility
    from money import NonNegativeMoney


class OrdineStorico(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    data: str
    file_sorgente: str = ""
    codice: str
    descrizione: str = ""
    quantita: int
    prezzo: NonNegativeMoney = 0
    batch_id: Optional[str] = None
    smart_venue_applicabile: bool = False
    caricato_smart_venue: Optional[bool] = None
    caricato_smart_venue_il: Optional[str] = None
    caricato_smart_venue_da: Optional[str] = None
    smart_venue_stato: str = "NON_APPLICABILE"
    smart_venue_movimento_id: Optional[str] = None
    smart_venue_riferimento_migrazione: Optional[str] = None
    operatore: str = ""
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    annullato: bool = False
    annullato_at: Optional[str] = None
    annullato_da: Optional[str] = None


class OrdineIn(BaseModel):
    data: str
    codice: str
    descrizione: Optional[str] = ""
    quantita: int
    prezzo: NonNegativeMoney = 0
    file_sorgente: Optional[str] = "manuale"
    idempotency_key: str


class BulkOrdineIn(BaseModel):
    file_sorgente: str = "carico manuale"
    data: Optional[str] = None
    idempotency_key: Optional[str] = None
    righe: List[Dict[str, Any]]


SMART_VENUE_AUTO_ORDER_REPAIR_ID = "smart-venue-auto-order-2026-10-12-v1"
SMART_VENUE_AUTO_ORDER_METADATA_MIGRATION_ID = "smart-venue-auto-order-2026-10-12-metadata-v1"
SMART_VENUE_AUTO_ORDER_REPAIR_EXPECTED = {
    "AMMS338": 150, "AMMS9": 60, "AMMS478": 40, "AMMS20155": 40,
    "AMMS20154": 40, "AMMS536": 60, "AMMS1784": 30, "AMMS20659": 60,
    "AMMS2982": 30, "AMMS3422": 20, "AMMS134": 20, "AMMS3906": 10,
    "AMMS3907": 30, "AMMS91": 30, "AMMS20041": 10, "AMMS20151": 20,
    "AMMS21037": 20, "AMMS3221": 20, "AMMS3458": 10, "AMMS951": 10,
    "AMMS81": 10, "AMMS177": 10, "AMMS3688": 20, "AMMS1733": 20,
    "AMMS20153": 20, "AMMS21460": 10, "AMMS245": 10,
}


class _OrderIdempotencyCollision(Exception):
    """Duplicate idempotency key detected after Mongo aborted a transaction."""


def _canonical_fingerprint(payload: Dict[str, Any]) -> str:
    canonical = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def parse_auto_order_excel(content: bytes, filename: str = "") -> Dict[str, Any]:
    """Parse the compact workbook exported by the Auto-Order page."""
    try:
        import openpyxl
    except Exception as exc:  # pragma: no cover - deployment dependency guard
        raise HTTPException(500, "openpyxl non installato") from exc

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content), data_only=True, read_only=True, keep_vba=False
        )
    except Exception as exc:
        raise HTTPException(422, f"File Excel non leggibile: {exc}") from exc

    if "Ordine fornitore" not in workbook.sheetnames:
        raise HTTPException(
            422,
            "Il file non è un Excel Auto-Order: manca il foglio 'Ordine fornitore'",
        )

    worksheet = workbook["Ordine fornitore"]
    header_row = None
    for row_index in range(1, min(worksheet.max_row, 20) + 1):
        labels = [
            str(worksheet.cell(row=row_index, column=column).value or "").strip().upper()
            for column in range(1, 5)
        ]
        if labels == ["CODICE", "ARTICOLO", "QTA", "TOTALE"]:
            header_row = row_index
            break
    if header_row is None:
        raise HTTPException(
            422,
            "Intestazioni Auto-Order non riconosciute: servono CODICE, ARTICOLO, QTA e TOTALE",
        )

    rows_by_code: Dict[str, Dict[str, Any]] = {}
    warnings: List[str] = []
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        raw_code = worksheet.cell(row=row_index, column=1).value
        raw_description = worksheet.cell(row=row_index, column=2).value
        raw_quantity = worksheet.cell(row=row_index, column=3).value
        raw_total = worksheet.cell(row=row_index, column=4).value
        code = str(raw_code or "").strip()
        if not code:
            continue
        if code.upper() == "TOTALE" or str(raw_quantity or "").strip().upper() == "TOTALE":
            continue
        try:
            numeric_quantity = float(raw_quantity)
            if not numeric_quantity.is_integer() or numeric_quantity <= 0:
                raise ValueError
            quantity = int(numeric_quantity)
        except (TypeError, ValueError):
            warnings.append(f"Riga {row_index}: quantità non valida per {code}, riga ignorata")
            continue
        try:
            row_total = float(raw_total or 0)
            if row_total < 0:
                raise ValueError
        except (TypeError, ValueError):
            warnings.append(f"Riga {row_index}: totale non valido per {code}, impostato a zero")
            row_total = 0

        description = str(raw_description or code).lstrip("'").strip() or code
        if code in rows_by_code:
            existing = rows_by_code[code]
            existing["quantita"] += quantity
            existing["totale_excel"] = round(existing["totale_excel"] + row_total, 2)
            existing["prezzo"] = round(existing["totale_excel"] / existing["quantita"], 2)
            warnings.append(f"Codice {code} ripetuto: quantità aggregate")
            continue
        rows_by_code[code] = {
            "codice": code,
            "descrizione": description,
            "quantita": quantity,
            "prezzo": round(row_total / quantity, 2) if quantity else 0,
            "totale_excel": round(row_total, 2),
        }

    rows = list(rows_by_code.values())
    if not rows:
        raise HTTPException(422, "Il file Auto-Order non contiene righe caricabili")

    raw_date = worksheet["D2"].value
    if hasattr(raw_date, "isoformat"):
        raw_date = raw_date.isoformat()
    return {
        "ok": True,
        "file": filename or "ordine_auto_order.xlsx",
        "fornitore": str(worksheet["B2"].value or "").lstrip("'").strip(),
        "data_ordine": str(raw_date or ""),
        "selezione": str(worksheet["B3"].value or "").lstrip("'").strip(),
        "righe": rows,
        "riepilogo": {
            "prodotti": len(rows),
            "pezzi": sum(row["quantita"] for row in rows),
            "totale": round(sum(row["totale_excel"] for row in rows), 2),
            "avvisi": len(warnings),
        },
        "avvisi": warnings,
    }


def receipt_quantities_by_code(rows: List[Dict[str, Any]]) -> Dict[str, int]:
    quantities: Dict[str, int] = {}
    for row in rows:
        code = str(row.get("codice") or "").strip()
        quantity = int(row.get("quantita", 0) or 0)
        if code and quantity > 0:
            quantities[code] = quantities.get(code, 0) + quantity
    return quantities


def create_orders_router(
    *,
    get_database,
    cap_limit,
    request_operator,
    session_kwargs,
    run_transactionally,
    create_backup_snapshot,
    product_model,
    categorize_description,
    normalize_nonnegative_money,
    record_audit_event,
    read_capped,
):
    router = APIRouter()

    def smart_state(applicable: bool, confirmed: bool = False) -> str:
        return "CONFERMATO" if confirmed else "DA_INSERIRE" if applicable else "NON_APPLICABILE"

    async def create_pending_smart_movement(*, order_id: str, batch_id: Optional[str],
                                            code: str, quantity: int, operator: str,
                                            created_at: str, session=None) -> str:
        """Create the external-load obligation. The unique key protects retry paths."""
        movement_id = f"smart-venue-load:{order_id}"
        document = {
            "id": movement_id, "idempotency_key": movement_id, "codice": code,
            "tipo": "CARICO_IN_ENTRATA", "quantita": quantity, "stato": "PENDENTE",
            "ordine_id": order_id, "batch_id": batch_id, "origine": "CARICO_MERCE",
            "operatore": operator, "created_at": created_at,
        }
        try:
            await get_database().smart_venue_movimenti.insert_one(document, **session_kwargs(session))
        except DuplicateKeyError:
            existing = await get_database().smart_venue_movimenti.find_one(
                {"idempotency_key": movement_id}, {"_id": 0}, **session_kwargs(session)
            )
            if not existing or existing.get("ordine_id") != order_id:
                raise HTTPException(409, "Movimento Smart Venue già presente ma non coerente")
        return movement_id

    @router.post("/ordini/import-auto-order/preview")
    async def preview_auto_order_excel(file: UploadFile = File(...)):
        filename = str(file.filename or "").replace("\\", "/").rsplit("/", 1)[-1][:180]
        if not filename.lower().endswith(".xlsx"):
            raise HTTPException(422, "Seleziona il file .xlsx esportato da Auto-Order")
        return parse_auto_order_excel(await read_capped(file), filename)

    async def repair_smart_venue_auto_order_2026_10_12() -> Dict[str, Any]:
        database = get_database()
        migration = await database.app_migrations.find_one(
            {"id": SMART_VENUE_AUTO_ORDER_REPAIR_ID}, {"_id": 0}
        )
        if migration:
            return {"ok": True, "gia_applicata": True, **migration}

        rows = await database.storico_ordini.find(
            {"file_sorgente": "da auto-order (modificato)"},
            {"_id": 0, "data": 1, "codice": 1, "quantita": 1},
        ).to_list(1000)
        rows = [
            row for row in rows
            if str(row.get("data") or "").startswith("2026-10-12")
        ]
        if not rows:
            return {"ok": True, "non_applicabile": True}
        actual = receipt_quantities_by_code(rows)
        if actual != SMART_VENUE_AUTO_ORDER_REPAIR_EXPECTED:
            raise RuntimeError(
                "Riparazione annullata: il carico storico del 12/10/2026 non "
                f"coincide con quello verificato (atteso {SMART_VENUE_AUTO_ORDER_REPAIR_EXPECTED}, "
                f"trovato {actual})."
            )

        codes = list(SMART_VENUE_AUTO_ORDER_REPAIR_EXPECTED)
        smart_rows = await database.smart_venue.find(
            {"codice": {"$in": codes}},
            {"_id": 0, "codice": 1, "riparazioni_applicate": 1},
        ).to_list(1000)
        by_code = {row["codice"]: row for row in smart_rows}
        missing = sorted(set(codes) - set(by_code))
        if missing:
            raise RuntimeError(
                "Riparazione annullata: mancano righe Smart Venue per "
                + ", ".join(missing)
            )

        backup = await create_backup_snapshot(
            "Prima riparazione Smart Venue carico Auto-Order 12/10/2026",
            "riparazione-smart-venue-carico",
            ["smart_venue", "storico_ordini"],
        )
        now = datetime.now(timezone.utc).isoformat()
        operations = [
            UpdateOne(
                {
                    "codice": code,
                    "riparazioni_applicate": {
                        "$ne": SMART_VENUE_AUTO_ORDER_REPAIR_ID
                    },
                },
                {
                    "$inc": {"smart_venue": quantity},
                    "$addToSet": {
                        "riparazioni_applicate": SMART_VENUE_AUTO_ORDER_REPAIR_ID
                    },
                    "$set": {"updated_at": now},
                },
            )
            for code, quantity in SMART_VENUE_AUTO_ORDER_REPAIR_EXPECTED.items()
        ]
        if operations:
            await database.smart_venue.bulk_write(operations, ordered=True)

        repaired = await database.smart_venue.count_documents({
            "codice": {"$in": codes},
            "riparazioni_applicate": SMART_VENUE_AUTO_ORDER_REPAIR_ID,
        })
        if repaired != len(codes):
            raise RuntimeError(
                f"Riparazione incompleta: marcate {repaired} righe Smart Venue su {len(codes)}."
            )
        migration = {
            "id": SMART_VENUE_AUTO_ORDER_REPAIR_ID,
            "created_at": now,
            "backup_id": backup["id"],
            "righe_riparate": len(codes),
            "pezzi_ripristinati": sum(
                SMART_VENUE_AUTO_ORDER_REPAIR_EXPECTED.values()
            ),
        }
        await database.app_migrations.insert_one(migration.copy())
        return {"ok": True, "gia_applicata": False, **migration}

    async def historical_smart_venue_metadata_migration(dry_run: bool = True) -> Dict[str, Any]:
        """Annotate the already repaired 12 October receipt, never touching balances."""
        database = get_database()
        migration = await database.app_migrations.find_one(
            {"id": SMART_VENUE_AUTO_ORDER_METADATA_MIGRATION_ID}, {"_id": 0}
        )
        if migration:
            return {"ok": True, "dry_run": dry_run, "gia_applicata": True, **migration}
        origin = await database.app_migrations.find_one(
            {"id": SMART_VENUE_AUTO_ORDER_REPAIR_ID}, {"_id": 0}
        )
        rows = await database.storico_ordini.find(
            {"file_sorgente": "da auto-order (modificato)"}, {"_id": 0}
        ).to_list(1000)
        rows = [row for row in rows if str(row.get("data") or "").startswith("2026-10-12")]
        actual = receipt_quantities_by_code(rows)
        anomalies = []
        if actual != SMART_VENUE_AUTO_ORDER_REPAIR_EXPECTED:
            anomalies.append({"tipo": "QUANTITA_STORICHE_NON_CORRISPONDENTI", "trovato": actual})
        if not origin or int(origin.get("pezzi_ripristinati") or 0) != 810:
            anomalies.append({"tipo": "RIPARAZIONE_ORIGINALE_NON_VERIFICATA"})
        codes = list(SMART_VENUE_AUTO_ORDER_REPAIR_EXPECTED)
        smart_rows = await database.smart_venue.find({"codice": {"$in": codes}}, {"_id": 0, "codice": 1, "riparazioni_applicate": 1}).to_list(1000)
        repaired = {row.get("codice") for row in smart_rows if SMART_VENUE_AUTO_ORDER_REPAIR_ID in (row.get("riparazioni_applicate") or [])}
        missing = sorted(set(codes) - repaired)
        if missing: anomalies.append({"tipo": "SALDI_ORIGINALI_NON_MARCATI", "codici": missing})
        report = {"ok": not anomalies, "dry_run": dry_run, "righe": len(rows), "pezzi": sum(actual.values()), "codici": sorted(actual), "anomalie": anomalies, "modifiche_previste": len(rows)}
        if dry_run or anomalies:
            return report
        backup = await create_backup_snapshot("Prima migrazione metadati Smart Venue 12/10/2026", "migrazione-smart-venue-metadata-2026-10-12", ["prodotti", "smart_venue", "storico_ordini", "ricezioni_merce", "rettifiche_inventario", "smart_venue_movimenti", "app_migrations"])
        now = datetime.now(timezone.utc).isoformat()
        result = await database.storico_ordini.update_many(
            {"file_sorgente": "da auto-order (modificato)", "data": {"$regex": "^2026-10-12"}, "codice": {"$in": codes}},
            {"$set": {"smart_venue_applicabile": True, "caricato_smart_venue": True,
                      "smart_venue_stato": "CONFERMATO_MIGRAZIONE_STORICA",
                      "caricato_smart_venue_il": now, "caricato_smart_venue_da": "MIGRAZIONE_STORICA",
                      "smart_venue_riferimento_migrazione": SMART_VENUE_AUTO_ORDER_REPAIR_ID}},
        )
        migration = {"id": SMART_VENUE_AUTO_ORDER_METADATA_MIGRATION_ID, "created_at": now,
                     "backup_id": backup["id"], "righe_aggiornate": result.modified_count,
                     "pezzi_verificati": 810, "riferimento_migrazione_origine": SMART_VENUE_AUTO_ORDER_REPAIR_ID}
        await database.app_migrations.insert_one(migration.copy())
        return {"ok": True, "dry_run": False, "gia_applicata": False, **migration}

    @router.get("/smart-venue/migrazioni/2026-10-12/dry-run")
    async def smart_venue_metadata_dry_run():
        return await historical_smart_venue_metadata_migration(dry_run=True)

    @router.post("/smart-venue/migrazioni/2026-10-12/applica")
    async def apply_smart_venue_metadata_migration():
        preview = await historical_smart_venue_metadata_migration(dry_run=True)
        if not preview["ok"]:
            raise HTTPException(409, preview)
        return await historical_smart_venue_metadata_migration(dry_run=False)

    @router.get("/ordini")
    async def list_ordini(limit: int = 1000):
        capped_limit = cap_limit(limit)
        return await get_database().storico_ordini.find(
            {"annullato": {"$ne": True}}, {"_id": 0}
        ).sort("data", -1).limit(capped_limit).to_list(capped_limit)

    @router.post("/ordini")
    async def add_ordine(order: OrdineIn, request: Request = None):
        if order.quantita <= 0:
            raise HTTPException(422, "La quantità deve essere maggiore di zero")
        key = str(order.idempotency_key or "").strip()
        if not 10 <= len(key) <= 160:
            raise HTTPException(422, "idempotency_key non valida")
        database = get_database()
        fingerprint_payload = order.model_dump(mode="json", exclude={"idempotency_key"})
        fingerprint_payload["codice"] = str(fingerprint_payload.get("codice") or "").strip()
        fingerprint = _canonical_fingerprint(fingerprint_payload)

        def replay(existing):
            if existing.get("request_fingerprint") != fingerprint:
                raise HTTPException(409, "Chiave idempotente già usata con dati differenti")
            return {**existing, "gia_registrato": True}

        existing = await database.storico_ordini.find_one(
            {"idempotency_key": key}, {"_id": 0}
        )
        if existing:
            return replay(existing)

        async def persist(session):
            kwargs = session_kwargs(session)
            smart_venue_applicable = bool(
                await database.smart_venue.find_one(
                    {"codice": order.codice}, {"_id": 1}, **kwargs
                )
            )
            document = OrdineStorico(
                **order.model_dump(exclude={"idempotency_key"}),
                smart_venue_applicabile=smart_venue_applicable,
                caricato_smart_venue=False if smart_venue_applicable else None,
                smart_venue_stato=smart_state(smart_venue_applicable),
                operatore=request_operator(request),
            ).model_dump()
            document.update({
                "idempotency_key": key,
                "request_fingerprint": fingerprint,
            })
            try:
                await database.storico_ordini.insert_one(document, **kwargs)
            except DuplicateKeyError as exc:
                raise _OrderIdempotencyCollision from exc

            product_changed = False
            movement_id = None
            try:
                if order.codice:
                    product_result = await database.prodotti.update_one(
                        {"codice": order.codice},
                        {"$inc": {
                            "giacenza_negozio": order.quantita,
                            "acquistati": order.quantita,
                        }},
                        **kwargs,
                    )
                    if product_result.matched_count != 1:
                        raise HTTPException(409, f"Prodotto {order.codice} non trovato")
                    product_changed = True
                    if smart_venue_applicable:
                        movement_id = await create_pending_smart_movement(
                            order_id=document["id"], batch_id=None,
                            code=order.codice, quantity=order.quantita,
                            operator=document["operatore"], created_at=document["created_at"],
                            session=session,
                        )
                        await database.storico_ordini.update_one(
                            {"id": document["id"]},
                            {"$set": {"smart_venue_movimento_id": movement_id}},
                            **kwargs,
                        )
                        document["smart_venue_movimento_id"] = movement_id
            except Exception:
                if session is None:
                    if movement_id:
                        await database.smart_venue_movimenti.delete_one(
                            {"id": movement_id}
                        )
                    if product_changed:
                        await database.prodotti.update_one(
                            {"codice": order.codice},
                            {"$inc": {
                                "giacenza_negozio": -order.quantita,
                                "acquistati": -order.quantita,
                            }},
                        )
                    await database.storico_ordini.delete_one({"id": document["id"]})
                raise
            return {**document, "gia_registrato": False}

        try:
            result = await run_transactionally(persist)
        except _OrderIdempotencyCollision:
            existing = await database.storico_ordini.find_one(
                {"idempotency_key": key}, {"_id": 0}
            )
            if existing:
                return replay(existing)
            raise HTTPException(409, "Ordine già in elaborazione")
        await record_audit_event(
            "CREATO", "ORDINE_MANUALE", result["id"],
            dopo=result, idempotency_key=key,
        )
        return result

    @router.post("/ordini/{order_id}/conferma-smart-venue")
    async def conferma_carico_smart_venue(
        order_id: str,
        request: Request = None,
    ):
        """Confirm a pending external load and increase the known external balance once."""
        database = get_database()
        current = await database.storico_ordini.find_one(
            {"id": order_id}, {"_id": 0}
        )
        if not current:
            raise HTTPException(404, "Movimento di carico non trovato")
        if not current.get("smart_venue_applicabile"):
            raise HTTPException(
                422,
                "Il prodotto di questo carico non è gestito in Smart Venue",
            )
        if current.get("caricato_smart_venue") is True:
            return {**current, "ok": True, "gia_confermato": True}
        if current.get("annullato"):
            raise HTTPException(409, "Il carico è annullato e non può essere confermato")

        now = datetime.now(timezone.utc).isoformat()
        async def confirm(session):
            kwargs = session_kwargs(session)
            movement = await database.smart_venue_movimenti.find_one({"ordine_id": order_id, "stato": "PENDENTE"}, {"_id": 0}, **kwargs)
            if not movement:
                raise HTTPException(409, "Movimento Smart Venue pendente non trovato")
            quantity = int(movement["quantita"])
            updated_smart = await database.smart_venue.update_one({"codice": current["codice"]}, {"$inc": {"smart_venue": quantity}, "$set": {"updated_at": now}}, **kwargs)
            if updated_smart.matched_count != 1: raise HTTPException(409, "Saldo Smart Venue cambiato durante la conferma")
            changed = await database.smart_venue_movimenti.update_one({"id": movement["id"], "stato": "PENDENTE"}, {"$set": {"stato": "CONFERMATO", "saldo_successivo": None, "confirmed_at": now, "operatore": request_operator(request)}}, **kwargs)
            if changed.matched_count != 1: raise HTTPException(409, "Movimento Smart Venue cambiato nel frattempo")
            return await database.storico_ordini.update_one({"id": order_id, "caricato_smart_venue": {"$ne": True}}, {"$set": {"caricato_smart_venue": True, "smart_venue_stato": "CONFERMATO", "caricato_smart_venue_il": now, "caricato_smart_venue_da": request_operator(request)}}, **kwargs)
        result = await run_transactionally(confirm)
        if result.modified_count != 1:
            latest = await database.storico_ordini.find_one(
                {"id": order_id}, {"_id": 0}
            )
            if latest and latest.get("caricato_smart_venue") is True:
                return {**latest, "ok": True, "gia_confermato": True}
            raise HTTPException(409, "Il movimento è cambiato nel frattempo")
        updated = await database.storico_ordini.find_one(
            {"id": order_id}, {"_id": 0}
        )
        return {**updated, "ok": True, "gia_confermato": False}

    @router.delete("/ordini/{order_id}")
    async def del_ordine(order_id: str, request: Request = None):
        """Annulla un carico manuale compensandone tutti gli effetti fisici.

        Lo storico rimane disponibile per audit, ma non viene più incluso nella
        lista operativa. I carichi bulk hanno il proprio ciclo batch e non
        possono essere annullati da questa route riga-per-riga.
        """
        database = get_database()

        async def annul(session):
            kwargs = session_kwargs(session)
            order = await database.storico_ordini.find_one(
                {"id": order_id}, {"_id": 0}, **kwargs
            )
            if not order:
                raise HTTPException(404, "Movimento di carico non trovato")
            if order.get("batch_id"):
                raise HTTPException(
                    409,
                    "I carichi batch non possono essere annullati riga per riga",
                )
            if order.get("annullato"):
                return {"ok": True, "gia_annullato": True, "id": order_id}
            if order.get("caricato_smart_venue") is True:
                raise HTTPException(
                    409,
                    "Carico già confermato in Smart Venue: annullamento bloccato finché non viene registrata una compensazione esterna esplicita",
                )

            code = str(order.get("codice") or "").strip()
            quantity = int(order.get("quantita") or 0)
            if not code or quantity <= 0:
                raise HTTPException(409, "Movimento di carico non compensabile")

            product = await database.prodotti.find_one({"codice": code}, {"_id": 1}, **kwargs)
            if not product:
                raise HTTPException(409, f"Prodotto {code} non trovato: annullamento bloccato")

            updated_product = await database.prodotti.update_one(
                {
                    "codice": code,
                    "giacenza_negozio": {"$gte": quantity},
                    "acquistati": {"$gte": quantity},
                },
                {"$inc": {"giacenza_negozio": -quantity, "acquistati": -quantity}},
                **kwargs,
            )
            if updated_product.matched_count != 1:
                raise HTTPException(
                    409,
                    "Il prodotto è cambiato: la giacenza o gli acquistati non consentono l'annullamento",
                )

            try:
                if order.get("smart_venue_applicabile"):
                    changed = await database.smart_venue_movimenti.update_one(
                        {"ordine_id": order_id, "stato": "PENDENTE"},
                        {"$set": {"stato": "ANNULLATO", "cancelled_at": datetime.now(timezone.utc).isoformat(), "operatore": request_operator(request)}},
                        **kwargs,
                    )
                    if changed.matched_count != 1:
                        raise HTTPException(409, "Movimento Smart Venue pendente non trovato o già modificato")

                now = datetime.now(timezone.utc).isoformat()
                marked = await database.storico_ordini.update_one(
                    {"id": order_id, "annullato": {"$ne": True}},
                    {"$set": {
                        "annullato": True,
                        "smart_venue_stato": "ANNULLATO",
                        "annullato_at": now,
                        "annullato_da": request_operator(request),
                    }},
                    **kwargs,
                )
                if marked.matched_count != 1:
                    raise HTTPException(409, "Movimento cambiato durante l'annullamento")
            except Exception:
                if session is None:
                    if order.get("smart_venue_applicabile"):
                        await database.smart_venue_movimenti.update_one(
                            {"ordine_id": order_id, "stato": "ANNULLATO"},
                            {"$set": {"stato": "PENDENTE"}},
                        )
                    await database.prodotti.update_one(
                        {"codice": code},
                        {"$inc": {"giacenza_negozio": quantity, "acquistati": quantity}},
                    )
                raise

            return {"ok": True, "gia_annullato": False, "id": order_id, "ordine": order}

        result = await run_transactionally(annul)
        if not result["gia_annullato"]:
            order = result.pop("ordine")
            await record_audit_event(
                "ANNULLATO",
                "ORDINE_MANUALE",
                order_id,
                prima=order,
                dopo={
                    "annullato": True,
                    "codice": order.get("codice"),
                    "quantita": order.get("quantita"),
                },
            )
        return result

    @router.post("/ordini/bulk")
    async def bulk_carico(body: BulkOrdineIn, request: Request = None):
        """Receive merchandise once and update stock and purchase counters."""
        database = get_database()
        data = body.data or datetime.now(timezone.utc).isoformat()
        idempotency_key = str(body.idempotency_key or uuid.uuid4()).strip()
        if len(idempotency_key) > 200:
            raise HTTPException(422, "Chiave idempotente troppo lunga")

        prepared = []
        errors = []
        for index, row in enumerate(body.righe):
            try:
                code = str(row.get("codice") or "").strip()
                numeric_quantity = float(row.get("quantita"))
                if not numeric_quantity.is_integer():
                    raise ValueError("la quantità deve essere intera")
                quantity = int(numeric_quantity)
                if not code and quantity == 0:
                    continue
                if not code:
                    raise ValueError("codice mancante")
                if quantity <= 0:
                    raise ValueError("la quantità deve essere maggiore di zero")
                price = normalize_nonnegative_money(row.get("prezzo", 0))
                if price < 0:
                    raise ValueError("il prezzo non può essere negativo")
                prepared.append({
                    "index": index,
                    "codice": code,
                    "quantita": quantity,
                    "prezzo": price,
                    "descrizione": str(row.get("descrizione") or ""),
                })
            except Exception as exc:
                errors.append({"riga": index + 1, "errore": str(exc)})
        if errors:
            raise HTTPException(
                422,
                {"message": "Correggi le righe non valide", "errori": errors},
            )
        if not prepared:
            raise HTTPException(422, "Nessuna riga valida da caricare")

        request_fingerprint = _canonical_fingerprint({
            "data": data,
            "file_sorgente": body.file_sorgente,
            "righe": [{
                "codice": row["codice"], "quantita": row["quantita"],
                "prezzo": str(row["prezzo"]), "descrizione": row["descrizione"],
            } for row in prepared],
        })

        def replay_batch(existing_batch):
            if existing_batch.get("request_fingerprint") != request_fingerprint:
                raise HTTPException(409, "Chiave idempotente già usata con dati differenti")
            if existing_batch.get("status") == "completed" and existing_batch.get("response"):
                return {**existing_batch["response"], "gia_registrato": True}
            raise HTTPException(409, "Questo carico è già in elaborazione: attendi e aggiorna la pagina")

        existing_batch = await database.ricezioni_merce.find_one(
            {"idempotency_key": idempotency_key}, {"_id": 0}
        )
        if existing_batch:
            return replay_batch(existing_batch)

        batch_id = str(uuid.uuid4())
        created_at = datetime.now(timezone.utc).isoformat()
        operator = request_operator(request)
        backup = await create_backup_snapshot(
            "Prima del carico merce",
            "pre-carico-merce",
            ["prodotti", "storico_ordini", "smart_venue", "ricezioni_merce"],
        )

        async def persist_receipt(session):
            kwargs = session_kwargs(session)
            await database.ricezioni_merce.insert_one({
                "id": batch_id,
                "idempotency_key": idempotency_key,
                "request_fingerprint": request_fingerprint,
                "status": "processing",
                "created_at": created_at,
                "data": data,
                "file_sorgente": body.file_sorgente,
                "operatore": operator,
            }, **kwargs)

            inserted = 0
            created_products = 0
            movements = []
            rollback_actions = []
            try:
                for row in prepared:
                    code = row["codice"]
                    quantity = row["quantita"]
                    product = (
                        await database.prodotti.find_one(
                            {"codice": code}, session=session
                        )
                        if session
                        else await database.prodotti.find_one({"codice": code})
                    )
                    smart_row = await database.smart_venue.find_one(
                        {"codice": code}, {"_id": 1}, **kwargs
                    )
                    smart_applicable = bool(smart_row)
                    movement = OrdineStorico(
                        id=f"{batch_id}:{row['index']}",
                        batch_id=batch_id,
                        data=data,
                        file_sorgente=body.file_sorgente,
                        codice=code,
                        descrizione=(
                            row["descrizione"]
                            or str((product or {}).get("descrizione") or code)
                        ),
                        quantita=quantity,
                        prezzo=row["prezzo"],
                        smart_venue_applicabile=smart_applicable,
                        caricato_smart_venue=False if smart_applicable else None,
                        smart_venue_stato=smart_state(smart_applicable),
                        operatore=operator,
                        created_at=created_at,
                    ).model_dump()
                    await database.storico_ordini.insert_one(dict(movement), **kwargs)
                    if smart_applicable:
                        movement_id = await create_pending_smart_movement(
                            order_id=movement["id"], batch_id=batch_id, code=code,
                            quantity=quantity, operator=operator, created_at=created_at,
                            session=session,
                        )
                        movement["smart_venue_movimento_id"] = movement_id
                        await database.storico_ordini.update_one(
                            {"id": movement["id"]},
                            {"$set": {"smart_venue_movimento_id": movement_id}}, **kwargs
                        )

                    if product:
                        update = await database.prodotti.update_one(
                            {"codice": code},
                            {"$inc": {
                                "giacenza_negozio": quantity,
                                "acquistati": quantity,
                            }},
                            **kwargs,
                        )
                        if update.matched_count != 1:
                            raise RuntimeError(
                                f"prodotto {code} cambiato durante il carico"
                            )
                        rollback_actions.append(
                            ("existing_product", code, quantity)
                        )
                    else:
                        new_product = product_model(
                            codice=code,
                            descrizione=row["descrizione"] or code,
                            categoria=categorize_description(
                                row["descrizione"], code
                            ),
                            prezzo=row["prezzo"],
                            acquistati=quantity,
                            giacenza_negozio=quantity,
                        ).model_dump()
                        await database.prodotti.insert_one(
                            dict(new_product), **kwargs
                        )
                        rollback_actions.append(
                            ("new_product", code, new_product["id"])
                        )
                        created_products += 1
                    inserted += 1
                    movements.append(movement)

                response = {
                    "caricate": inserted,
                    "prodotti_nuovi": created_products,
                    "errori": [],
                    "batch_id": batch_id,
                    "idempotency_key": idempotency_key,
                    "movimenti": movements,
                    "gia_registrato": False,
                    "backup_id": backup["id"],
                }
                completed = await database.ricezioni_merce.update_one(
                    {"id": batch_id, "status": "processing"},
                    {"$set": {
                        "status": "completed",
                        "completed_at": datetime.now(timezone.utc).isoformat(),
                        "response": response,
                    }},
                    **kwargs,
                )
                if completed.matched_count != 1:
                    raise RuntimeError(
                        "stato del carico cambiato durante il salvataggio"
                    )
                return response
            except Exception:
                if session is None:
                    for kind, code, value in reversed(rollback_actions):
                        if kind == "existing_product":
                            await database.prodotti.update_one(
                                {"codice": code},
                                {"$inc": {
                                    "giacenza_negozio": -value,
                                    "acquistati": -value,
                                }},
                            )
                        elif kind == "new_product":
                            await database.prodotti.delete_one(
                                {"codice": code, "id": value}
                            )
                    await database.storico_ordini.delete_many(
                        {"batch_id": batch_id}
                    )
                    await database.smart_venue_movimenti.delete_many({"batch_id": batch_id})
                    await database.ricezioni_merce.update_one(
                        {"id": batch_id},
                        {"$set": {
                            "status": "failed",
                            "failed_at": datetime.now(timezone.utc).isoformat(),
                        }},
                    )
                raise

        try:
            return await run_transactionally(persist_receipt)
        except DuplicateKeyError:
            existing_batch = await database.ricezioni_merce.find_one(
                {"idempotency_key": idempotency_key}, {"_id": 0}
            )
            if existing_batch:
                return replay_batch(existing_batch)
            raise HTTPException(409, "Questo carico è già stato inviato")

    return {
        "router": router,
        "list_ordini": list_ordini,
        "add_ordine": add_ordine,
        "conferma_carico_smart_venue": conferma_carico_smart_venue,
        "del_ordine": del_ordine,
        "bulk_carico": bulk_carico,
        "preview_auto_order_excel": preview_auto_order_excel,
        "repair_smart_venue_auto_order_2026_10_12": repair_smart_venue_auto_order_2026_10_12,
        "historical_smart_venue_metadata_migration": historical_smart_venue_metadata_migration,
    }
